# -*- coding: utf-8 -*-
"""Extracción robusta de máxima demanda de cabecera desde Excel de mediciones.

Flujo (preciso, sin ambigüedad):
  1) Lookup alimentador → medidor + Vll en ``medidoralimentador.xlsx``
     (ignora filas Barra/LT; exige medidor no vacío).
  2) Abrir archivo de ``medicioncabecera`` (hoja = código medidor, match flexible).
     Si el archivo elegido no tiene la hoja, auto-localiza el correcto.
  3) Fila de máximo ``kW tot`` (empate → última fecha); kvar/kVA de esa misma fila.
  4) P promedio y factor de carga = P_prom / P_max × 100.
  5) Validaciones físicas (P>0, Q presente, S≈√(P²+Q²) con aviso).
"""
from __future__ import print_function

import math
import os
import re
import threading
from datetime import datetime

from core.clientes_suministro import resolve_dir

_CACHE_LOCK = threading.Lock()
_MAP_CACHE = {"mtime": None, "path": None, "by_feeder": {}}
_SHEET_INDEX_CACHE = {}  # path -> {mtime, sheets_norm: {norm: real_name}}


def medicioncabecera_dir(settings=None):
    return resolve_dir(settings, "medicioncabecera_dir", "data/input/common/medicioncabecera")


def medidoralimentador_dir(settings=None):
    return resolve_dir(settings, "medidoralimentador_dir", "data/input/common/medidoralimentador")


def medidoralimentador_path(settings=None):
    d = medidoralimentador_dir(settings)
    preferred = (settings or {}).get("medidoralimentador_file") or "medidoralimentador.xlsx"
    cand = os.path.join(d, preferred)
    if os.path.isfile(cand):
        return cand
    if not os.path.isdir(d):
        return cand
    for name in sorted(os.listdir(d)):
        low = name.lower()
        if name.startswith("~$"):
            continue
        if low.endswith((".xlsx", ".xlsm")):
            return os.path.join(d, name)
    return cand


def list_medicioncabecera_files(settings=None):
    d = medicioncabecera_dir(settings)
    if not os.path.isdir(d):
        return []
    out = []
    for name in sorted(os.listdir(d)):
        if name.startswith("~$"):
            continue
        low = name.lower()
        if low.endswith((".xls", ".xlsx", ".xlsm")):
            path = os.path.join(d, name)
            out.append({
                "name": name,
                "path": path,
                "size_mb": round(os.path.getsize(path) / (1024.0 * 1024.0), 1),
            })
    return out


def _norm_feeder(fid):
    s = str(fid or "").strip().upper()
    s = s.replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", "", s)


def _norm_meter(code):
    s = str(code or "").strip().upper()
    s = s.replace("–", "-").replace("—", "-")
    s = re.sub(r"\s+", "", s)
    return s


def _meter_aliases(code):
    """Variantes de nombre de hoja posibles para un medidor."""
    want = _norm_meter(code)
    if not want:
        return []
    aliases = [want]
    if want.startswith("L-"):
        aliases.append(want[2:])
    elif want.startswith("L") and len(want) > 1 and want[1].isdigit():
        aliases.append(want[1:])
    else:
        aliases.append("L-" + want)
    # sin guiones
    aliases.append(want.replace("-", ""))
    if want.startswith("L-"):
        aliases.append(("L" + want[2:]).replace("-", ""))
    # únicos preservando orden
    seen = set()
    out = []
    for a in aliases:
        if a and a not in seen:
            seen.add(a)
            out.append(a)
    return out


def _to_float(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        try:
            f = float(v)
        except Exception:
            return None
        if math.isnan(f) or math.isinf(f):
            return None
        return f
    s = str(v).strip().replace("\xa0", " ").replace(" ", "")
    if not s or s.lower() in ("nan", "none", "-", "#n/a", "#nv", "#value!"):
        return None
    s = s.replace(",", ".")
    s = re.sub(r"[^\d.\-eE]", "", s)
    if not s or s in (".", "-", "-.", "e", "E"):
        return None
    try:
        f = float(s)
    except Exception:
        return None
    if math.isnan(f) or math.isinf(f):
        return None
    return f


def _normalize_vll(v):
    n = _to_float(v)
    if n is None:
        return None
    # Valores típicos Electro Dunas; también aceptar 22900 V → 22.9 kV
    if n > 100:  # probablemente voltios
        n = n / 1000.0
    if abs(n - 22.9) < 0.08:
        return 22.9
    if abs(n - 10.0) < 0.08:
        return 10.0
    if abs(n - 60.0) < 0.5:
        return 60.0
    return round(n, 3)


def _is_real_feeder_code(code):
    """True si parece alimentador radial (PAxxx, INxxx…), no Barra/LT genérico."""
    c = _norm_feeder(code)
    if not c:
        return False
    if c in ("LT", "BARRA", "BARRAIND", "BARRA IND"):
        return False
    if c.startswith("BARRA"):
        return False
    # típico: 2 letras + dígitos (PA217, SI112, IN112…)
    if re.match(r"^[A-Z]{1,3}\d{2,4}[A-Z]?$", c):
        return True
    # permitir códigos con guión tipo PA-217
    if re.match(r"^[A-Z]{1,3}-\d{2,4}$", c):
        return True
    return False


def _load_medidor_map(settings=None, force=False):
    """Carga/caché del mapa alimentador → medidor/Vll."""
    path = medidoralimentador_path(settings)
    if not os.path.isfile(path):
        raise RuntimeError("No existe medidoralimentador: %s" % path)
    try:
        mtime = os.path.getmtime(path)
    except Exception:
        mtime = None

    with _CACHE_LOCK:
        if (
            not force
            and _MAP_CACHE["path"] == path
            and _MAP_CACHE["mtime"] == mtime
            and _MAP_CACHE["by_feeder"]
        ):
            return _MAP_CACHE["by_feeder"], path

    from openpyxl import load_workbook

    by_feeder = {}
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        headers = None
        col_alim = col_med = col_v = col_sig = 0
        for row in ws.iter_rows(values_only=True):
            if not row or all(c in (None, "") for c in row):
                continue
            cells = [("" if c is None else str(c).strip()) for c in row]
            lower = [c.lower() for c in cells]
            if headers is None and any(
                ("alimentador" in h) or (h == "medidor") for h in lower
            ):
                headers = lower
                col_alim = next(
                    (i for i, h in enumerate(headers) if "alimentador" in h or h in ("codigo", "código")),
                    0,
                )
                col_med = next((i for i, h in enumerate(headers) if "medidor" in h), 1)
                col_v = next(
                    (
                        i
                        for i, h in enumerate(headers)
                        if ("tension" in h) or ("tensión" in h) or h in ("vll", "kv", "nivel")
                    ),
                    2,
                )
                col_sig = next((i for i, h in enumerate(headers) if "sigla" in h), 3)
                continue
            if headers is None:
                continue
            raw_code = row[col_alim] if col_alim < len(row) else ""
            code = _norm_feeder(raw_code)
            if not code:
                continue
            medidor_raw = row[col_med] if col_med < len(row) else ""
            medidor = str(medidor_raw or "").strip()
            if not medidor:
                continue
            vll = _normalize_vll(row[col_v] if col_v < len(row) else None)
            siglas = ""
            if col_sig < len(row) and row[col_sig] not in (None, ""):
                siglas = str(row[col_sig]).strip()
            entry = {
                "feeder_id": code,
                "feeder_raw": str(raw_code or "").strip(),
                "medidor": medidor,
                "Vll_kV": vll,
                "siglas": siglas,
                "is_feeder": _is_real_feeder_code(code),
            }
            # Preferir filas de alimentador real si hay colisión (LT/Barra)
            prev = by_feeder.get(code)
            if prev is None or (entry["is_feeder"] and not prev.get("is_feeder")):
                by_feeder[code] = entry
    finally:
        wb.close()

    with _CACHE_LOCK:
        _MAP_CACHE["path"] = path
        _MAP_CACHE["mtime"] = mtime
        _MAP_CACHE["by_feeder"] = by_feeder
    return by_feeder, path


def _feeder_aliases(feeder_id):
    """Variantes BD ↔ medidor (IN112↔SI112, etc.)."""
    fid = _norm_feeder(feeder_id)
    if not fid:
        return []
    out = [fid]
    m = re.match(r"^([A-Z]{1,3})(\d{2,4}[A-Z]?)$", fid)
    if m:
        pref, num = m.group(1), m.group(2)
        swap = {"IN": "SI", "SI": "IN"}
        if pref in swap:
            out.append(swap[pref] + num)
    seen = set()
    uniq = []
    for a in out:
        if a and a not in seen:
            seen.add(a)
            uniq.append(a)
    return uniq


def _preferred_sistema_keywords(siglas, feeder_id):
    key = (siglas or "").strip().upper()
    if not key:
        m = re.match(r"^([A-Z]{1,3})", _norm_feeder(feeder_id) or "")
        key = m.group(1) if m else ""
    if key == "IN":
        key = "SI"
    return {
        "SI": ["ica", "pisco"],
        "PA": ["pisco", "ica"],
        "PI": ["pisco"],
        "AL": ["chincha"],
        "NA": ["nasca"],
        "CH": ["chincha"],
        "PO": ["nasca", "pisco"],
    }.get(key, [])


def lookup_feeder_medidor(feeder_id, settings=None):
    """Devuelve {feeder_id, medidor, Vll_kV, siglas} desde medidoralimentador.xlsx."""
    fid = _norm_feeder(feeder_id)
    if not fid:
        raise RuntimeError("Indique código de alimentador")

    by_feeder, path = _load_medidor_map(settings=settings)
    entry = None
    matched_as = None
    for alias in _feeder_aliases(fid):
        if alias in by_feeder:
            entry = by_feeder[alias]
            matched_as = alias
            break
        bare = alias.replace("-", "")
        for k, v in by_feeder.items():
            if k.replace("-", "") == bare:
                entry = v
                matched_as = k
                break
        if entry is not None:
            break

    if entry is None:
        suggestions = []
        for a in _feeder_aliases(fid):
            if a != fid:
                suggestions.append(a)
        prefix = re.sub(r"\d+.*$", "", fid)
        suggestions.extend(
            sorted(
                k for k in by_feeder.keys()
                if k.startswith(prefix) and by_feeder[k].get("is_feeder")
            )[:6]
        )
        hint = (" · pruebe: " + ", ".join(suggestions[:8])) if suggestions else ""
        raise RuntimeError(
            "Alimentador %s no está en medidoralimentador (%s)%s"
            % (fid, os.path.basename(path), hint)
        )

    if not entry.get("medidor"):
        raise RuntimeError("Alimentador %s sin medidor en %s" % (fid, os.path.basename(path)))

    vll = entry.get("Vll_kV")
    if vll is None:
        raise RuntimeError(
            "Alimentador %s sin nivel de tensión en medidoralimentador" % fid
        )

    return {
        "feeder_id": fid,
        "map_code": matched_as or entry.get("feeder_id") or fid,
        "medidor": entry["medidor"],
        "Vll_kV": vll,
        "siglas": entry.get("siglas") or "",
        "source": path,
        "is_feeder": bool(entry.get("is_feeder")),
    }


def _require_xlrd():
    try:
        import xlrd  # noqa: F401
        return xlrd
    except ImportError:
        raise RuntimeError(
            "Falta xlrd==1.2.0 en Python37 (.tools). "
            "Instale con: .tools\\python37-win32\\python.exe -m pip install xlrd==1.2.0"
        )


def _sheet_index_for_file(xls_path):
    """Índice norm→nombre real de hojas, cacheado por mtime."""
    xlrd = _require_xlrd()
    path = os.path.abspath(xls_path)
    try:
        mtime = os.path.getmtime(path)
    except Exception:
        mtime = None
    with _CACHE_LOCK:
        cached = _SHEET_INDEX_CACHE.get(path)
        if cached and cached.get("mtime") == mtime and cached.get("sheets_norm"):
            return cached["sheets_norm"], cached.get("sheet_names") or []

    book = xlrd.open_workbook(path, on_demand=True)
    try:
        names = list(book.sheet_names())
        sheets_norm = {}
        for n in names:
            if not n or str(n).startswith("|") or str(n) == "Presentation Sheet":
                continue
            key = _norm_meter(n)
            if key and key not in sheets_norm:
                sheets_norm[key] = n
            # también sin guiones
            bare = key.replace("-", "")
            if bare and bare not in sheets_norm:
                sheets_norm[bare] = n
    finally:
        book.release_resources()

    with _CACHE_LOCK:
        _SHEET_INDEX_CACHE[path] = {
            "mtime": mtime,
            "sheets_norm": sheets_norm,
            "sheet_names": names,
        }
    return sheets_norm, names


def _find_sheet_name(sheet_names_or_index, medidor):
    """Resuelve nombre de hoja real para un medidor (match flexible, preciso)."""
    if isinstance(sheet_names_or_index, dict):
        index = sheet_names_or_index
    else:
        index = {}
        for n in sheet_names_or_index or []:
            key = _norm_meter(n)
            if key and key not in index:
                index[key] = n
            bare = key.replace("-", "")
            if bare and bare not in index:
                index[bare] = n

    for alias in _meter_aliases(medidor):
        if alias in index:
            return index[alias]
        bare = alias.replace("-", "")
        if bare in index:
            return index[bare]
    return None


def _col_index_strict(headers, exact_names, contains_names=None):
    """Primero match exacto (casefold), luego contains ordenado por candidatos."""
    lower = [str(h or "").strip().lower() for h in headers]
    exact_set = [e.lower() for e in exact_names]
    for i, h in enumerate(lower):
        if h in exact_set:
            return i
    # normalizar espacios múltiples
    compact = [re.sub(r"\s+", " ", h) for h in lower]
    for e in exact_set:
        e2 = re.sub(r"\s+", " ", e)
        for i, h in enumerate(compact):
            if h == e2:
                return i
    for cand in contains_names or ():
        c = cand.lower()
        for i, h in enumerate(lower):
            if c == h or (len(c) >= 3 and c in h):
                return i
    return None


def _cell_datetime(book, sheet, row, col):
    xlrd = _require_xlrd()
    cell = sheet.cell(row, col)
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate_as_datetime(cell.value, book.datemode)
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        try:
            return xlrd.xldate_as_datetime(cell.value, book.datemode)
        except Exception:
            pass
    v = cell.value
    if isinstance(v, datetime):
        return v
    s = str(v or "").strip()
    if not s:
        return None
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%m/%d/%Y %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%d-%m-%Y %H:%M:%S",
        "%d-%m-%Y %H:%M",
    ):
        try:
            return datetime.strptime(s, fmt)
        except Exception:
            continue
    return None


def _fmt_fecha(dt):
    if dt is None:
        return ""
    if isinstance(dt, datetime):
        return "%d/%02d/%04d %02d:%02d:%02d" % (
            dt.day, dt.month, dt.year, dt.hour, dt.minute, dt.second
        )
    return str(dt)


def _row_sort_key_time(book, sheet, row, i_time):
    if i_time is None:
        return datetime.min
    dt = _cell_datetime(book, sheet, row, i_time)
    if isinstance(dt, datetime):
        return dt
    return datetime.min


def _sheet_metric_kind(headers):
    """Clasifica hoja: 'full' | 'kw' | 'kvar' | 'kva' | None."""
    i_kw = _col_index_strict(headers, exact_names=("kW tot", "kw tot"), contains_names=("kw tot",))
    i_kvar = _col_index_strict(headers, exact_names=("kVAR tot", "kvar tot"), contains_names=("kvar tot",))
    i_kva = _col_index_strict(headers, exact_names=("kVA tot", "kva tot"), contains_names=("kva tot",))
    if i_kw is not None and i_kvar is not None:
        return "full", i_kw, i_kvar, i_kva
    if i_kw is not None and i_kvar is None and i_kva is None:
        return "kw", i_kw, None, None
    if i_kvar is not None and i_kw is None:
        return "kvar", None, i_kvar, None
    if i_kva is not None and i_kw is None:
        return "kva", None, None, i_kva
    return None, i_kw, i_kvar, i_kva


def _sheets_for_feeder(all_names, feeder_id, medidor):
    """Hojas candidatas: medidor exacto + familia IN112 / IN112V18 / …"""
    codes = []
    for a in _feeder_aliases(feeder_id) + _meter_aliases(medidor) + [_norm_feeder(feeder_id)]:
        if a and a not in codes:
            codes.append(a)
            codes.append(a.replace("-", ""))
    codes = [c for c in codes if c]
    hits = []
    for n in all_names or []:
        if not n or str(n).startswith("|") or str(n) == "Presentation Sheet":
            continue
        nu = _norm_meter(n)
        nub = nu.replace("-", "")
        for c in codes:
            cb = c.replace("-", "")
            if nu == c or nub == cb or nub.startswith(cb):
                hits.append(n)
                break
    # Preferir nombres cortos / V18 (kW) primero
    def rank(name):
        u = _norm_meter(name).replace("-", "")
        score = 0
        if u in codes or u.replace("L", "", 1) in codes:
            score += 100
        if u.endswith("V18") or u.endswith("18"):
            score += 40
        if "_V19" in name.upper() or u.endswith("V19"):
            score += 10
        if "_20" in name.upper() or u.endswith("20"):
            score += 10
        return (-score, len(name), name)
    return sorted(set(hits), key=rank)


def _value_at_time(book, sheet, i_time, i_val, target_dt, data_start=1):
    """Busca valor en la misma Local Time (tolerancia 1s)."""
    if target_dt is None or i_time is None or i_val is None:
        return None
    best = None
    best_delta = None
    for r in range(data_start, sheet.nrows):
        dt = _cell_datetime(book, sheet, r, i_time)
        if not isinstance(dt, datetime):
            continue
        delta = abs((dt - target_dt).total_seconds())
        if delta <= 1.0:
            v = _to_float(sheet.cell_value(r, i_val))
            if v is not None:
                return v
        if best_delta is None or delta < best_delta:
            best_delta = delta
            best = _to_float(sheet.cell_value(r, i_val))
    if best_delta is not None and best_delta <= 900:  # 15 min
        return best
    return None


def _finalize_stats(max_kw, kvar, kva, fecha_dt, sum_kw, n, n_skip, sheet_name, medidor, xls_path, row_max, warnings):
    if max_kw is None or not n or max_kw <= 0:
        raise RuntimeError("Sin Pmáx válida en %s" % sheet_name)
    if kvar is None:
        raise RuntimeError("Sin kvar en la fecha de Pmáx (%s)" % sheet_name)
    if kva is None:
        kva = math.sqrt(max_kw * max_kw + kvar * kvar)
        warnings.append("S_kVA derivado de √(P²+Q²)")
    else:
        s_calc = math.sqrt(max_kw * max_kw + kvar * kvar)
        if s_calc > 1e-6 and abs(kva - s_calc) / s_calc > 0.15:
            warnings.append(
                "S medido (%.2f) difiere de √(P²+Q²)=%.2f" % (kva, s_calc)
            )
    p_avg = sum_kw / float(n)
    fdc = p_avg / max_kw * 100.0
    if fdc < 5.0 or fdc > 100.5:
        warnings.append("factor de carga fuera de rango típico (%.1f%%)" % fdc)
    if fecha_dt is None:
        warnings.append("fecha de Pmáx no disponible")
    return {
        "sheet": sheet_name,
        "medidor": medidor,
        "file": os.path.basename(xls_path),
        "path": xls_path,
        "P_kW": round(max_kw, 2),
        "Q_kvar": round(kvar, 2),
        "S_kVA": round(kva, 2),
        "P_avg_kW": round(p_avg, 2),
        "factor_carga_pct": round(fdc, 2),
        "fecha_medicion": _fmt_fecha(fecha_dt),
        "n_samples": n,
        "n_skipped": n_skip,
        "row_max": row_max,
        "warnings": warnings,
    }


def extract_max_demanda_from_xls(xls_path, medidor, feeder_id=None):
    """Lee hoja(s) del medidor/alimentador y extrae P_max, Q/kVA, P_prom y FdC.

    Soporta:
      - hoja única con kW+kvar(+kVA)  (p.ej. L-PA235)
      - hojas partidas tipo Ica (IN112V18=kW, IN112_V19=kvar, IN112_20=kVA)
    """
    xlrd = _require_xlrd()
    if not os.path.isfile(xls_path):
        raise RuntimeError("No existe archivo de medición: %s" % xls_path)

    sheets_norm, all_names = _sheet_index_for_file(xls_path)
    candidates = _sheets_for_feeder(all_names, feeder_id or "", medidor)
    sheet_exact = _find_sheet_name(sheets_norm, medidor)
    if sheet_exact and sheet_exact not in candidates:
        candidates.insert(0, sheet_exact)
    if not candidates:
        raise RuntimeError(
            "Medidor/alimentador %s/%s sin hoja en %s"
            % (medidor, feeder_id or "?", os.path.basename(xls_path))
        )

    book = xlrd.open_workbook(xls_path, on_demand=True)
    try:
        # 1) Preferir hoja completa (kW+kvar juntos)
        for sheet_name in candidates:
            sh = book.sheet_by_name(sheet_name)
            if sh.nrows < 2:
                continue
            headers = [sh.cell_value(0, c) for c in range(sh.ncols)]
            kind, i_kw, i_kvar, i_kva = _sheet_metric_kind(headers)
            if kind != "full":
                continue
            i_time = _col_index_strict(
                headers,
                exact_names=("Local Time", "Timestamp"),
                contains_names=("local time", "timestamp"),
            )
            max_kw = None
            max_row = None
            sum_kw = 0.0
            n = 0
            n_skip = 0
            for r in range(1, sh.nrows):
                kw = _to_float(sh.cell_value(r, i_kw))
                if kw is None or kw < -1.0 or kw > 5.0e5:
                    n_skip += 1
                    continue
                sum_kw += kw
                n += 1
                if max_kw is None or kw > max_kw + 1e-9:
                    max_kw, max_row = kw, r
                elif abs(kw - max_kw) <= 1e-9 and i_time is not None:
                    if _row_sort_key_time(book, sh, r, i_time) >= _row_sort_key_time(
                        book, sh, max_row, i_time
                    ):
                        max_row = r
            if max_row is None:
                continue
            kvar = _to_float(sh.cell_value(max_row, i_kvar))
            kva = _to_float(sh.cell_value(max_row, i_kva)) if i_kva is not None else None
            fecha_dt = _cell_datetime(book, sh, max_row, i_time) if i_time is not None else None
            return _finalize_stats(
                max_kw, kvar, kva, fecha_dt, sum_kw, n, n_skip,
                sheet_name, medidor, xls_path, max_row + 1, [],
            )

        # 2) Hojas partidas (Ica): ensamblar kW + kvar + kVA por fecha
        kw_sheet = kvar_sheet = kva_sheet = None
        kw_icol = kvar_icol = kva_icol = None
        for sheet_name in candidates:
            sh = book.sheet_by_name(sheet_name)
            if sh.nrows < 2:
                continue
            headers = [sh.cell_value(0, c) for c in range(sh.ncols)]
            kind, i_kw, i_kvar, i_kva = _sheet_metric_kind(headers)
            if kind == "kw" and kw_sheet is None:
                kw_sheet, kw_icol = sheet_name, i_kw
            elif kind == "kvar" and kvar_sheet is None:
                kvar_sheet, kvar_icol = sheet_name, i_kvar
            elif kind == "kva" and kva_sheet is None:
                kva_sheet, kva_icol = sheet_name, i_kva
        if kw_sheet is None:
            raise RuntimeError(
                "No hay hoja con 'kW tot' para %s en %s (candidatas: %s)"
                % (feeder_id or medidor, os.path.basename(xls_path), ", ".join(candidates[:8]))
            )
        if kvar_sheet is None:
            raise RuntimeError(
                "No hay hoja con 'kVAR tot' para %s en %s (hay kW en %s)"
                % (feeder_id or medidor, os.path.basename(xls_path), kw_sheet)
            )

        sh_kw = book.sheet_by_name(kw_sheet)
        hdr_kw = [sh_kw.cell_value(0, c) for c in range(sh_kw.ncols)]
        i_time = _col_index_strict(
            hdr_kw,
            exact_names=("Local Time", "Timestamp"),
            contains_names=("local time", "timestamp"),
        )
        max_kw = None
        max_row = None
        sum_kw = 0.0
        n = 0
        n_skip = 0
        for r in range(1, sh_kw.nrows):
            kw = _to_float(sh_kw.cell_value(r, kw_icol))
            if kw is None or kw < -1.0 or kw > 5.0e5:
                n_skip += 1
                continue
            sum_kw += kw
            n += 1
            if max_kw is None or kw > max_kw + 1e-9:
                max_kw, max_row = kw, r
            elif abs(kw - max_kw) <= 1e-9 and i_time is not None:
                if _row_sort_key_time(book, sh_kw, r, i_time) >= _row_sort_key_time(
                    book, sh_kw, max_row, i_time
                ):
                    max_row = r
        if max_row is None:
            raise RuntimeError("Hoja %s sin kW válidos" % kw_sheet)
        fecha_dt = _cell_datetime(book, sh_kw, max_row, i_time) if i_time is not None else None

        sh_q = book.sheet_by_name(kvar_sheet)
        hdr_q = [sh_q.cell_value(0, c) for c in range(sh_q.ncols)]
        i_time_q = _col_index_strict(
            hdr_q, exact_names=("Local Time", "Timestamp"), contains_names=("local time", "timestamp")
        )
        kvar = _value_at_time(book, sh_q, i_time_q, kvar_icol, fecha_dt)
        if kvar is None and max_row < sh_q.nrows:
            kvar = _to_float(sh_q.cell_value(max_row, kvar_icol))

        kva = None
        if kva_sheet is not None:
            sh_s = book.sheet_by_name(kva_sheet)
            hdr_s = [sh_s.cell_value(0, c) for c in range(sh_s.ncols)]
            i_time_s = _col_index_strict(
                hdr_s, exact_names=("Local Time", "Timestamp"), contains_names=("local time", "timestamp")
            )
            kva = _value_at_time(book, sh_s, i_time_s, kva_icol, fecha_dt)
            if kva is None and max_row < sh_s.nrows:
                kva = _to_float(sh_s.cell_value(max_row, kva_icol))

        warnings = [
            "Hojas partidas: kW=%s · kvar=%s%s"
            % (kw_sheet, kvar_sheet, (" · kVA=" + kva_sheet) if kva_sheet else "")
        ]
        return _finalize_stats(
            max_kw, kvar, kva, fecha_dt, sum_kw, n, n_skip,
            kw_sheet, medidor, xls_path, max_row + 1, warnings,
        )
    finally:
        book.release_resources()


def _file_has_feeder_data(path, medidor, feeder_id):
    try:
        sheets_norm, names = _sheet_index_for_file(path)
        if _find_sheet_name(sheets_norm, medidor):
            return True, _find_sheet_name(sheets_norm, medidor)
        cands = _sheets_for_feeder(names, feeder_id or "", medidor)
        if not cands:
            return False, None
        # verificar que al menos una tenga kW
        xlrd = _require_xlrd()
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            for sn in cands:
                sh = book.sheet_by_name(sn)
                if sh.nrows < 2:
                    continue
                headers = [sh.cell_value(0, c) for c in range(sh.ncols)]
                kind, _, _, _ = _sheet_metric_kind(headers)
                if kind in ("full", "kw"):
                    return True, sn
        finally:
            book.release_resources()
        return False, None
    except Exception:
        return False, None


def find_medicion_files_for_medidor(medidor, settings=None, files=None, feeder_id=None, siglas=None):
    """Lista Excel con hoja del medidor o familia de hojas del alimentador."""
    items = files if files is not None else list_medicioncabecera_files(settings)
    hits = []
    for item in items:
        path = item.get("path") if isinstance(item, dict) else item
        name = item.get("name") if isinstance(item, dict) else os.path.basename(path or "")
        if not path or not os.path.isfile(path):
            continue
        ok, sheet = _file_has_feeder_data(path, medidor, feeder_id)
        if ok:
            hits.append({
                "name": name,
                "path": path,
                "sheet": sheet,
                "size_mb": round(os.path.getsize(path) / (1024.0 * 1024.0), 1),
            })
    # Ordenar por sistema preferido
    prefs = _preferred_sistema_keywords(siglas, feeder_id or "")
    if prefs:
        def rank(h):
            low = (h.get("name") or "").lower()
            for i, p in enumerate(prefs):
                if p in low:
                    return i
            return 99
        hits.sort(key=rank)
    return hits


def find_medicion_file_for_medidor(medidor, settings=None, files=None, feeder_id=None, siglas=None):
    hits = find_medicion_files_for_medidor(
        medidor, settings=settings, files=files, feeder_id=feeder_id, siglas=siglas
    )
    if not hits:
        return None
    return hits[0]["path"]


def resolve_medicion_path(file_name_or_path, settings=None):
    if not file_name_or_path:
        return None
    raw = str(file_name_or_path).strip()
    if os.path.isfile(raw):
        return os.path.abspath(raw)
    d = medicioncabecera_dir(settings)
    cand = os.path.join(d, os.path.basename(raw))
    if os.path.isfile(cand):
        return os.path.abspath(cand)
    if os.path.isdir(d):
        want = os.path.basename(raw).lower()
        for name in os.listdir(d):
            if name.lower() == want and not name.startswith("~$"):
                return os.path.abspath(os.path.join(d, name))
    return None


def resolve_cabecera_medicion(feeder_id, settings=None):
    """Solo lookup: medidor, Vll y archivos Excel candidatos."""
    meta = lookup_feeder_medidor(feeder_id, settings=settings)
    hits = find_medicion_files_for_medidor(
        meta["medidor"],
        settings=settings,
        feeder_id=meta["feeder_id"],
        siglas=meta.get("siglas"),
    )
    return {
        "ok": True,
        "feeder_id": meta["feeder_id"],
        "map_code": meta.get("map_code"),
        "medidor": meta["medidor"],
        "Vll_kV": meta["Vll_kV"],
        "siglas": meta.get("siglas") or "",
        "mapping_source": meta.get("source"),
        "candidate_files": hits,
        "n_candidates": len(hits),
        "suggested_file": hits[0]["name"] if hits else None,
        "msg": (
            "Medidor %s (mapa %s) · Vll %s kV · %d archivo(s)"
            % (meta["medidor"], meta.get("map_code") or meta["feeder_id"], meta["Vll_kV"], len(hits))
        ),
    }


def extract_cabecera_medicion(feeder_id, medicion_file=None, settings=None, auto_find_file=True):
    """Pipeline completo: mapping + extracción (con alias IN↔SI y hojas partidas)."""
    meta = lookup_feeder_medidor(feeder_id, settings=settings)
    warnings = []
    file_switched = False
    requested = medicion_file
    if meta.get("map_code") and meta["map_code"] != meta["feeder_id"]:
        warnings.append("Mapa medidor: %s -> %s" % (meta["feeder_id"], meta["map_code"]))

    path = resolve_medicion_path(medicion_file, settings=settings) if medicion_file else None
    if path and not os.path.isfile(path):
        path = None

    def _has_data(p):
        ok, _ = _file_has_feeder_data(p, meta["medidor"], meta["feeder_id"])
        return ok

    if path and not _has_data(path):
        if auto_find_file:
            alt = find_medicion_file_for_medidor(
                meta["medidor"],
                settings=settings,
                feeder_id=meta["feeder_id"],
                siglas=meta.get("siglas"),
            )
            if alt:
                warnings.append(
                    "Archivo '%s' sin datos de %s; se usó '%s'"
                    % (os.path.basename(path), meta["feeder_id"], os.path.basename(alt))
                )
                path = alt
                file_switched = True
            else:
                raise RuntimeError(
                    "Sin medición de %s/%s en '%s' ni en otros Excel"
                    % (meta["feeder_id"], meta["medidor"], os.path.basename(path))
                )
        else:
            raise RuntimeError(
                "Sin hoja de %s/%s en %s"
                % (meta["feeder_id"], meta["medidor"], os.path.basename(path))
            )

    if not path and auto_find_file:
        path = find_medicion_file_for_medidor(
            meta["medidor"],
            settings=settings,
            feeder_id=meta["feeder_id"],
            siglas=meta.get("siglas"),
        )
        if path:
            file_switched = True
            warnings.append("Archivo auto-seleccionado: %s" % os.path.basename(path))

    if not path:
        raise RuntimeError(
            "No hay Excel en medicioncabecera con datos de %s (medidor %s)"
            % (meta["feeder_id"], meta["medidor"])
        )

    stats = extract_max_demanda_from_xls(
        path, meta["medidor"], feeder_id=meta["feeder_id"]
    )
    warnings.extend(stats.get("warnings") or [])

    out = {
        "ok": True,
        "feeder_id": meta["feeder_id"],
        "map_code": meta.get("map_code"),
        "medidor": meta["medidor"],
        "Vll_kV": meta["Vll_kV"],
        "siglas": meta.get("siglas") or "",
        "mapping_source": meta.get("source"),
        "P_kW": stats["P_kW"],
        "Q_kvar": stats["Q_kvar"],
        "S_kVA": stats["S_kVA"],
        "P_avg_kW": stats["P_avg_kW"],
        "factor_carga_pct": stats["factor_carga_pct"],
        "fecha_medicion": stats["fecha_medicion"],
        "medicion_file": stats["file"],
        "medicion_path": stats["path"],
        "requested_file": os.path.basename(str(requested)) if requested else None,
        "file_switched": file_switched,
        "sheet": stats["sheet"],
        "n_samples": stats["n_samples"],
        "n_skipped": stats.get("n_skipped") or 0,
        "row_max": stats.get("row_max"),
        "warnings": warnings,
        "msg": (
            "%s · medidor %s · Máx P=%.2f · Q=%.2f · S=%.2f · Pprom=%.2f · FdC=%.1f%% · %s · %s"
            % (
                meta["feeder_id"],
                meta["medidor"],
                stats["P_kW"] or 0,
                stats["Q_kvar"] or 0,
                stats["S_kVA"] or 0,
                stats["P_avg_kW"] or 0,
                stats["factor_carga_pct"] or 0,
                stats["fecha_medicion"] or "",
                stats["file"],
            )
        ),
    }
    if warnings:
        out["msg"] = out["msg"] + " · avisos: " + "; ".join(warnings[:2])
    return out
