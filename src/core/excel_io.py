from __future__ import print_function

def _xl():
    try:
        import openpyxl
        return openpyxl
    except Exception:
        raise RuntimeError("Instale openpyxl en el Python37 de CYME usando scripts\\00_install_dependencies.bat")

def read_rows(path, sheet):
    ox=_xl()
    wb=ox.load_workbook(path,data_only=True,read_only=True)
    if sheet not in wb.sheetnames:
        raise RuntimeError("No existe hoja: "+sheet)
    ws=wb[sheet]
    it=ws.iter_rows(values_only=True)
    # Saltar titulo y fila vacía, cabecera en fila 3
    next(it); next(it)
    headers=[str(x).strip() if x is not None else "" for x in next(it)]
    out=[]
    for vals in it:
        row={h:v for h,v in zip(headers,vals) if h}
        if any(v not in (None,"") for v in row.values()):
            out.append(row)
    return out

def read_kv(path, sheet, key_col="Parametro", val_col="Valor"):
    rows=read_rows(path,sheet)
    return {str(r.get(key_col,"")).strip():r.get(val_col) for r in rows if r.get(key_col)}

def write_kv(path, sheet, updates, key_col="Parametro", val_col="Valor"):
    """Actualiza celdas Valor en hoja tipo Parametro/Valor (misma estructura que read_kv)."""
    ox = _xl()
    wb = ox.load_workbook(path)
    if sheet not in wb.sheetnames:
        raise RuntimeError("No existe hoja: " + sheet)
    ws = wb[sheet]
    # Fila 3 = cabeceras (como read_rows)
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[3]]
    try:
        ki = headers.index(key_col)
        vi = headers.index(val_col)
    except ValueError:
        raise RuntimeError("Hoja %s sin columnas %s/%s" % (sheet, key_col, val_col))
    wanted = {str(k).strip(): v for k, v in (updates or {}).items()}
    found = set()
    for row in ws.iter_rows(min_row=4):
        key = row[ki].value
        if key is None:
            continue
        key = str(key).strip()
        if key in wanted:
            row[vi].value = wanted[key]
            found.add(key)
    missing = [k for k in wanted if k not in found]
    if missing:
        # Append nuevas filas
        for k in missing:
            ws.append([None] * len(headers))
            r = ws.max_row
            ws.cell(r, ki + 1, k)
            ws.cell(r, vi + 1, wanted[k])
    wb.save(path)
    return sorted(found)
