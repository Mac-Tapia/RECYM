# -*- coding: utf-8 -*-
"""
Sincroniza parámetros de equipos desde Excel (input) hacia la biblioteca CYMDIST.
Solo crea/actualiza equipos; no inventa IDs si ya existen equivalentes en la BD.
Basado en tutorial CYME EquipmentModeling (eq.Add / SetValue).
"""
from __future__ import print_function
import os
from openpyxl import load_workbook
from core.common import require_cympy, load_json, write_csv, run_cympy_main
from core.cympy_adapter import CymPyAdapter
from core.feeder_context import load_settings, output_path, catalog_path
from core.equipment_library import (
    inventory_library, save_inventory, equipment_exists, SIZE_HINTS,
)

def _num(v, default=None):
    if v is None or v == "":
        return default
    try:
        return float(str(v).replace(",", "."))
    except Exception:
        return default

def _read_aaac_excel(path):
    """Filas de calibre desde AAAC_CYMDIST / Lineas_Aereas."""
    rows = []
    wb = load_workbook(path, data_only=True, read_only=True)
    # Prefer sheet AAAC_CYMDIST
    sheet = "AAAC_CYMDIST" if "AAAC_CYMDIST" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet]
    header = None
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        vals = list(row)
        if not any(vals):
            continue
        first = str(vals[0] or "")
        if "Calibre" in first or first.strip() in ("35", "50", "70", "120"):
            if "Calibre" in first:
                header = vals
                continue
            # data row without waiting header if numeric first col
            size = _num(vals[0])
            if size is None:
                continue
            rows.append({
                "size_mm2": size,
                "strands": _num(vals[2]),
                "od_mm": _num(vals[4]),
                "od_cm": _num(vals[5]),
                "gmr_cm": _num(vals[8]),
                "ampacity": _num(vals[9]),
                "rdc20": _num(vals[10]),
                "r25": _num(vals[11]),
                # cols: ... R80, Icc Nexans, Icc captura CYMDIST
                "icc_nexans": _num(vals[13]) if len(vals) > 13 else None,
            })
    wb.close()
    return rows

def _read_cable_excel(path):
    rows = []
    wb = load_workbook(path, data_only=True, read_only=True)
    sheet = "CYMDIST_Carga_Rapida" if "CYMDIST_Carga_Rapida" in wb.sheetnames else None
    if not sheet:
        wb.close()
        return rows
    ws = wb[sheet]
    started = False
    for row in ws.iter_rows(values_only=True):
        vals = list(row)
        if not vals or vals[0] is None:
            continue
        first = str(vals[0])
        if "Sección" in first or "Seccion" in first:
            started = True
            continue
        if not started:
            # also accept bare numeric
            if _num(vals[0]) is None:
                continue
        size = _num(vals[0])
        if size is None:
            continue
        # CYMDIST_Carga_Rapida columns from preview: size, strands, type, material, voltage, R?, ...
        rows.append({
            "size_mm2": size,
            "strands": _num(vals[1]),
            "r1": _num(vals[7]) or _num(vals[5]),
            "x1": _num(vals[8]) or _num(vals[6]),
            "ampacity": _num(vals[10]) or _num(vals[9]),
        })
    wb.close()
    return rows

def _read_catalog_lines(path):
    """Catalogo_Maestro!Lineas_Aereas -> size hints."""
    rows = []
    if not os.path.isfile(path):
        return rows
    wb = load_workbook(path, data_only=True, read_only=True)
    if "Lineas_Aereas" not in wb.sheetnames:
        wb.close()
        return rows
    ws = wb["Lineas_Aereas"]
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 3:
            continue
        vals = list(row)
        if not vals or not vals[0]:
            continue
        eid = str(vals[0]).strip()
        size = _num(vals[3])
        rows.append({"catalog_id": eid, "size_mm2": size, "r25": _num(vals[5]), "gmr": _num(vals[8] if len(vals) > 8 else None)})
    wb.close()
    return rows

def _safe_set(eq_obj, prop, value):
    if value is None:
        return False
    try:
        eq_obj.SetValue(float(value) if isinstance(value, (int, float)) or str(value).replace(".", "", 1).replace(",", "", 1).isdigit() else value, prop)
        return True
    except Exception:
        try:
            eq_obj.SetValue(float(value), prop)
            return True
        except Exception:
            return False

def _safe_set_any(eq_obj, props, value):
    """Prueba varios nombres de propiedad CYMDIST hasta que uno acepte el valor."""
    for prop in props:
        if _safe_set(eq_obj, prop, value):
            return prop
    return None

# Nombres alternos según versión CYME / captura Electro Dunas
_COND_R_PROPS = (
    "Resistance", "Resistance25C", "DCResistance", "R25", "Rdc",
    "PositiveSequenceResistance",
)
_COND_AMP_PROPS = ("Ampacity", "ContinuousCurrent", "RatedCurrent", "NominalAmpacity")
_COND_ICC_PROPS = ("WithstandCurrent", "ShortCircuitCurrent", "FirstCycleWithstand")

def sync_conductors_from_aaac(cympy, aaac_rows, report):
    """Actualiza conductores AAAC### con GMR/diámetro/hilos/R/ampacidad del Excel fabricante."""
    et = cympy.enums.EquipmentType.Conductor
    for r in aaac_rows:
        size = int(r["size_mm2"]) if r.get("size_mm2") else None
        if not size:
            continue
        cid = SIZE_HINTS["Conductor"].get(size)
        if not cid or not equipment_exists(cympy, cid, "Conductor"):
            report.append({"Accion": "skip_conductor", "ID": cid or "", "Motivo": "no existe en BD", "Size": size})
            continue
        e = cympy.eq.GetEquipment(cid, et)
        changed = []
        if r.get("gmr_cm") is not None and _safe_set(e, "GMR", r["gmr_cm"]):
            changed.append("GMR")
        od = r.get("od_cm")
        if od is None and r.get("od_mm") is not None:
            od = float(r["od_mm"]) / 10.0
        if od is not None and _safe_set(e, "OutsideDiameter", od):
            changed.append("OutsideDiameter")
        if r.get("strands") is not None and _safe_set(e, "NumberOfStrands", r["strands"]):
            changed.append("NumberOfStrands")
        if r.get("size_mm2") is not None and _safe_set(e, "Size", r["size_mm2"]):
            changed.append("Size")
        # Preferir R@25°C (cymR25C); si no, Rdc@20°C de ficha
        r_ohm = r.get("r25") if r.get("r25") is not None else r.get("rdc20")
        prop_r = _safe_set_any(e, _COND_R_PROPS, r_ohm) if r_ohm is not None else None
        if prop_r:
            changed.append(prop_r)
        prop_a = _safe_set_any(e, _COND_AMP_PROPS, r.get("ampacity")) if r.get("ampacity") is not None else None
        if prop_a:
            changed.append(prop_a)
        icc = r.get("icc_nexans") or r.get("icc")
        prop_i = _safe_set_any(e, _COND_ICC_PROPS, icc) if icc is not None else None
        if prop_i:
            changed.append(prop_i)
        report.append({"Accion": "update_conductor", "ID": cid, "Campos": ",".join(changed) or "(sin cambio)", "Size": size})
        print("CONDUCTOR", cid, changed)

def ensure_oh_lines(cympy, aaac_rows, report, spacing_id="ATVB1"):
    """
    Asegura OverheadLine ATVB1-22.9KV-XXX por cada calibre del Excel,
    reutilizando conductores AAAC existentes (manual EquipmentModeling).
    Tambien alinea R1 (R@25) y ratings de ampacidad desde ficha fabricante.
    """
    et = cympy.enums.EquipmentType.OverheadLine
    for r in aaac_rows:
        size = int(r["size_mm2"]) if r.get("size_mm2") else None
        if not size:
            continue
        line_id = SIZE_HINTS["OverheadLine"].get(size)
        cond_id = SIZE_HINTS["Conductor"].get(size)
        if not line_id:
            continue
        if equipment_exists(cympy, line_id, "OverheadLine"):
            e = cympy.eq.GetEquipment(line_id, et)
            changed = []
            if cond_id and equipment_exists(cympy, cond_id, "Conductor"):
                if _safe_set(e, "PhaseConductorID", cond_id):
                    changed.append("PhaseConductorID")
            # En Electro Dunas R/ampacidad viven en OverheadLineDB, no en ConductorDB
            r_ohm = r.get("r25") if r.get("r25") is not None else r.get("rdc20")
            if r_ohm is not None and _safe_set(e, "PositiveSequenceResistance", r_ohm):
                changed.append("PositiveSequenceResistance")
            amp = r.get("ampacity")
            if amp is not None:
                for prop in ("NominalRating", "FirstRating", "SecondRating", "ThirdRating", "FourthRating"):
                    if _safe_set(e, prop, amp):
                        changed.append(prop)
            report.append({
                "Accion": "oh_update" if changed else "oh_exists",
                "ID": line_id,
                "Conductor": cond_id or "",
                "Campos": ",".join(changed),
                "Size": size,
            })
            if changed:
                print("OH", line_id, changed)
            continue
        # Crear solo si hay conductor base
        if not cond_id or not equipment_exists(cympy, cond_id, "Conductor"):
            report.append({"Accion": "oh_skip_create", "ID": line_id or "", "Motivo": "falta conductor", "Size": size})
            continue
        try:
            e = cympy.eq.Add(line_id, et)
            _safe_set(e, "PhaseConductorID", cond_id)
            if equipment_exists(cympy, "NONE", "Conductor"):
                _safe_set(e, "NeutralConductorID", "NONE")
            if spacing_id and equipment_exists(cympy, spacing_id, "OverheadSpacingOfConductor"):
                try:
                    e.SetValue(spacing_id, "SpacingID")
                except Exception:
                    pass
            r_ohm = r.get("r25") if r.get("r25") is not None else r.get("rdc20")
            if r_ohm is not None:
                _safe_set(e, "PositiveSequenceResistance", r_ohm)
            amp = r.get("ampacity")
            if amp is not None:
                for prop in ("NominalRating", "FirstRating", "SecondRating", "ThirdRating", "FourthRating"):
                    _safe_set(e, prop, amp)
            report.append({"Accion": "oh_created", "ID": line_id, "Conductor": cond_id, "Size": size})
            print("CREATED OH", line_id, "->", cond_id)
        except Exception as ex:
            report.append({"Accion": "oh_create_error", "ID": line_id, "Motivo": str(ex), "Size": size})

def ensure_cables(cympy, cable_rows, report):
    et = cympy.enums.EquipmentType.Cable
    for r in cable_rows:
        size = int(r["size_mm2"]) if r.get("size_mm2") else None
        if not size:
            continue
        cid = SIZE_HINTS["Cable"].get(size)
        if not cid:
            # no inventar nombres fuera del mapa conocido
            report.append({"Accion": "cable_skip", "ID": "", "Motivo": "sin mapeo SIZE_HINTS", "Size": size})
            continue
        if equipment_exists(cympy, cid, "Cable"):
            e = cympy.eq.GetEquipment(cid, et)
            changed = []
            if r.get("r1") is not None and _safe_set(e, "PositiveSequenceResistance", r["r1"]):
                changed.append("R1")
            if r.get("x1") is not None and _safe_set(e, "PositiveSequenceReactance", r["x1"]):
                changed.append("X1")
            report.append({"Accion": "cable_update", "ID": cid, "Campos": ",".join(changed) or "(ok)", "Size": size})
            print("CABLE", cid, changed)
            continue
        try:
            e = cympy.eq.Add(cid, et)
            if r.get("r1") is not None:
                _safe_set(e, "PositiveSequenceResistance", r["r1"])
            if r.get("x1") is not None:
                _safe_set(e, "PositiveSequenceReactance", r["x1"])
            report.append({"Accion": "cable_created", "ID": cid, "Size": size})
            print("CREATED CABLE", cid)
        except Exception as ex:
            report.append({"Accion": "cable_create_error", "ID": cid, "Motivo": str(ex), "Size": size})

def main():
    s = load_settings()
    api = load_json("config/cympy_api_map.json")
    c = require_cympy(s)
    a = CymPyAdapter(c, api, s)
    a.open_study()

    eq_dir = s.get("common_equipment_dir") or "data/input/common/equipment"
    if not os.path.isabs(eq_dir):
        root = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
        eq_dir = os.path.join(root, eq_dir)

    aaac_path = os.path.join(eq_dir, "AAAC_6201_T81_CYMDIST_9_referenciado.xlsx")
    cable_path = os.path.join(eq_dir, "Cable_Subterraneo_Cobre_XLPE_18-30kV_CYMDIST_v4.xlsx")

    report = []
    aaac_rows = _read_aaac_excel(aaac_path) if os.path.isfile(aaac_path) else []
    cable_rows = _read_cable_excel(cable_path) if os.path.isfile(cable_path) else []
    print("Excel AAAC filas:", len(aaac_rows), "| Cable filas:", len(cable_rows))

    # Inventario antes
    inv_before = inventory_library(c)
    save_inventory(output_path(s, "inventory", "equipment_library_before.json"), inv_before)

    if not s.get("dry_run"):
        sync_conductors_from_aaac(c, aaac_rows, report)
        ensure_oh_lines(c, aaac_rows, report)
        ensure_cables(c, cable_rows, report)
        # Guardar BD de equipos vía study save (equipos viven en MDB conectada)
        try:
            a.save_study()
        except Exception as ex:
            print("AVISO save:", ex)
    else:
        print("DRY_RUN sync equipment")

    inv_after = inventory_library(c)
    save_inventory(output_path(s, "inventory", "equipment_library.json"), inv_after)
    out = output_path(s, "inventory", "equipment_sync_report.csv")
    write_csv(out, report, ["Accion", "ID", "Conductor", "Campos", "Motivo", "Size"])
    print("Biblioteca OverheadLine:", inv_after.get("OverheadLine"))
    print("Biblioteca Cable:", inv_after.get("Cable"))
    print("Biblioteca Conductor:", inv_after.get("Conductor"))
    print(out)

if __name__ == "__main__":
    run_cympy_main(main)
