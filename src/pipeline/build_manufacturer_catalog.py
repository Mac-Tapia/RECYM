# -*- coding: utf-8 -*-
"""Build the global normalized manufacturer equipment catalog."""
from __future__ import print_function

import json
import os
from datetime import datetime

from core.common import p

CATALOG_VERSION = 1
OUTPUT_PATH = p("config", "manufacturer_equipment_catalog.json")

SOURCES = {
    "Cable": {
        "file": "data/input/common/equipment/Cable_Subterraneo_Cobre_XLPE_18-30kV_CYMDIST_v4.xlsx",
        "sheet": "Cu_XLPE_18-30kV",
        "manufacturer": "Prysmian",
        "family": "Cu XLPE 18/30(36) kV",
        "url": "https://asean.prysmian.com/sites/asean.prysmian.com/files/media/documents/MV%20Cables%20-%20250219.pdf",
    },
    "OverheadLine": {
        "file": "data/input/common/equipment/AAAC_6201_T81_CYMDIST_9_referenciado.xlsx",
        "sheet": "AAAC_CYMDIST",
        "manufacturer": "Nexans/Prysmian reference",
        "family": "AAAC 6201-T81",
        "url": "https://www.nexans.co/es/products/Redes-de-Transmisi%C3%B3n-y-Distribuci%C3%B3n-de-Energ%C3%ADa/Cables-de-Aluminio-Desnudo/Cable-AAAC/AAAC-6201-T81.html",
    },
}


def _number(value):
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", "."))
    except Exception:
        return None


def _header_map(row):
    return {
        str(value or "").replace("\n", " ").strip().lower(): index
        for index, value in enumerate(row)
        if value not in (None, "")
    }


def _value(row, headers, *names):
    for name in names:
        index = headers.get(name.lower())
        if index is not None and index < len(row):
            return row[index]
    return None


def _load_source(source):
    from openpyxl import load_workbook

    path = p(*source["file"].split("/"))
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook[source["sheet"]]
    rows = sheet.iter_rows(values_only=True)
    title = next(rows, ())
    header_row = ()
    for candidate in rows:
        first = str(candidate[0] or "").lower() if candidate else ""
        if "sección" in first or "calibre" in first:
            header_row = candidate
            break
    headers = _header_map(header_row)
    records = []
    for row in rows:
        section = _number(_value(row, headers, "sección nominal (mm²)", "calibre nominal (mm²)"))
        if section is None:
            continue
        record = {
            "equipment_family": source["family"],
            "equipment_type": source["type"],
            "manufacturer": source["manufacturer"],
            "nominal_section_mm2": section,
            "real_section_mm2": _number(_value(row, headers, "sección real (mm²)")),
            "material": _value(row, headers, "material"),
            "voltage_uo_u_kv": _value(row, headers, "tensión uo/u (kv)", "tensión (kv)"),
            "strands": _number(_value(row, headers, "n.º hilos", "cantidad de hilos (mín. iec 60228 clase 2)")),
            "rdc20_ohm_km": _number(_value(row, headers, "rdc @20°c (ω/km)", "rdc20_ohm_km")),
            "rdc25_ohm_km": _number(_value(row, headers, "rdc @25°c* (ω/km)", "rdc25_ohm_km", "r @25°c* (ω/km)")),
            "rdc90_ohm_km": _number(_value(row, headers, "rdc @90°c* (ω/km)", "rdc90_ohm_km", "r @90°c* (ω/km)")),
            "ampacity_air_a": _number(_value(row, headers, "ampacidad (a)", "ampacidad aire (a)", "amp aire (a)")),
            "ampacity_buried_a": _number(_value(row, headers, "ampacidad enterrado (a)", "amp. enterrado (a)")),
            "short_circuit_1s_ka": _number(_value(row, headers, "icc 1 s (ka)", "icc 1s (ka)", "icc 1 s (a)", "icc 1s (a)")),
            "source_file": source["file"],
            "source_sheet": source["sheet"],
            "source_url": source["url"],
            "source_title": str(title[0]) if title else "",
            "status": "manufacturer_reference",
        }
        records.append(record)
    return records


def build_catalog(output_path=OUTPUT_PATH):
    records = []
    for equipment_type, source in SOURCES.items():
        source = dict(source)
        source["type"] = equipment_type
        records.extend(_load_source(source))
    payload = {
        "schema_version": CATALOG_VERSION,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "scope": "global",
        "replacement_policy": {
            "only_default": True,
            "require_section_mm2": True,
            "preserve_non_default": True,
            "preserve_operating_voltage": True,
        },
        "records": records,
    }
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as handle:
        json.dump(payload, handle, indent=2, ensure_ascii=False)
    return payload


if __name__ == "__main__":
    result = build_catalog()
    print("Catalogo global: %s registros -> %s" % (len(result["records"]), OUTPUT_PATH))
