# -*- coding: utf-8 -*-
from __future__ import print_function
"""Genera Correcciones propuestas desde el diagnostico CYMDIST.

Usa el catálogo oficial config/diagnostic_corrections_catalog.json
(cymdist.cymsg [DiagnosticTool] 220000–220053 + LoadFlow/Simulation).

TODO Error/Warning/Hint del NetworkDiagnostic genera fila en la tabla
Correcciones; las de auto_apply=true se aplican en bulk_fix.
"""
import csv
import os
from core.common import load_json, write_csv, truthy, require_cympy, run_cympy_main
from core.feeder_context import load_settings, output_path, catalog_path
from core.cympy_adapter import CymPyAdapter
from core.equipment_library import inventory_library, save_inventory
from pipeline.diagnostic_registry import propose_row, load_catalog

FIX_SEVERITIES = ("Error", "Warning", "Hint")


def _read_diag_csv(path):
    if not os.path.isfile(path):
        raise RuntimeError("No existe diagnostico. Ejecute primero run_network_diagnostic.py: " + path)
    with open(path, "r", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def _catalog_size_hints(s):
    hints = {}
    path = catalog_path(s)
    if not os.path.isfile(path):
        return hints
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        if "Lineas_Aereas" in wb.sheetnames:
            for i, row in enumerate(wb["Lineas_Aereas"].iter_rows(values_only=True)):
                if i < 3 or not row or not row[0]:
                    continue
                try:
                    hints[("OverheadLine", str(row[0]).strip())] = (
                        float(row[3]) if row[3] not in (None, "") else None
                    )
                except Exception:
                    pass
        if "Cables_Subterraneos" in wb.sheetnames:
            for i, row in enumerate(wb["Cables_Subterraneos"].iter_rows(values_only=True)):
                if i < 3 or not row or not row[0]:
                    continue
                try:
                    hints[("Cable", str(row[0]).strip())] = (
                        float(row[4]) if len(row) > 4 and row[4] not in (None, "") else None
                    )
                except Exception:
                    pass
        wb.close()
    except Exception as ex:
        print("AVISO catalogo size hints:", ex)
    return hints


def _row_needs_fix(r):
    sev = (r.get("Severidad") or "").strip()
    flag = (r.get("Requiere_Correccion") or "").strip().upper()
    if flag == "NO":
        return False
    if flag == "SI":
        return True
    return sev in FIX_SEVERITIES


def main():
    s = load_settings()
    cat = load_catalog()
    print("[%s] Catalogo diagnostico: %s codigos" % (s["feeder_id"], len(cat)))

    diag_after = output_path(s, "diagnostics", "cymdist_diagnostic_errors_after.csv")
    diag = output_path(s, "diagnostics", "cymdist_diagnostic_errors.csv")
    if os.path.isfile(diag_after):
        rows_after = _read_diag_csv(diag_after)
        n_prob = sum(1 for r in rows_after if _row_needs_fix(r))
        if n_prob == 0:
            print("Usando diagnostico AFTER (sin problemas):", diag_after)
            rows = rows_after
            diag = diag_after
        else:
            rows = _read_diag_csv(diag)
    else:
        rows = _read_diag_csv(diag)

    inv = {}
    cympy = None
    try:
        api = load_json("config/cympy_api_map.json")
        cympy = require_cympy(s)
        a = CymPyAdapter(cympy, api, s)
        a.open_study()
        inv = inventory_library(cympy)
        save_inventory(output_path(s, "inventory", "equipment_library.json"), inv)
    except Exception as ex:
        print("AVISO: sin CYMDIST para validar equipos:", ex)

    ctx = {
        "cympy": cympy,
        "inv": inv,
        "defaults": s.get("default_equipment") or {},
        "catalog_hints": _catalog_size_hints(s),
        "vbase": s.get("voltage_ll_kv") or 22.9,
    }

    proposed = []
    seen = set()
    for r in rows:
        if not _row_needs_fix(r):
            continue
        prop = propose_row(r, s, ctx)
        key = (
            prop.get("Codigo"),
            prop.get("Tipo"),
            prop.get("ID_CYMDIST"),
            prop.get("Accion_Sugerida"),
        )
        if key in seen:
            continue
        seen.add(key)
        proposed.append(prop)

    out = output_path(s, "diagnostics", "correcciones_propuestas.csv")
    headers = [
        "Activo", "Tipo", "ID_CYMDIST", "ID_Seccion", "Equipo_Actual", "Equipo_Nuevo",
        "BaseVoltage_kV", "Fase", "Observacion", "Origen_Dato", "Codigo",
        "Accion_Sugerida", "Severidad", "Mensaje_CYME",
    ]
    write_csv(out, proposed, headers)
    activos = sum(1 for p in proposed if truthy(p.get("Activo")))
    by_action = {}
    for p in proposed:
        a = p.get("Accion_Sugerida") or "?"
        by_action[a] = by_action.get(a, 0) + 1
    print("[%s] Propuestas: %s (activas=%s) desde %s" % (
        s["feeder_id"], len(proposed), activos, diag))
    print("Por accion:", by_action)
    print(out)

    if s.get("sync_corrections_to_catalog"):
        _sync_catalog(s, [p for p in proposed if truthy(p.get("Activo"))])


def _sync_catalog(s, rows):
    try:
        import openpyxl
    except Exception as ex:
        print("No se pudo sync catalogo:", ex)
        return
    path = catalog_path(s)
    if not os.path.isfile(path):
        print("Sin catalogo:", path)
        return
    wb = openpyxl.load_workbook(path)
    if "Correcciones" not in wb.sheetnames:
        print("Sin hoja Correcciones")
        return
    ws = wb["Correcciones"]
    if ws.max_row >= 4:
        ws.delete_rows(4, ws.max_row - 3)
    headers = [c.value for c in ws[3]]
    for r in rows:
        line = []
        for h in headers:
            if h == "Activo":
                line.append(True)
            elif h in r:
                line.append(r.get(h))
            else:
                line.append(None)
        ws.append(line)
    wb.save(path)
    print("Catalogo actualizado:", path, "filas", len(rows))


if __name__ == "__main__":
    run_cympy_main(main)
