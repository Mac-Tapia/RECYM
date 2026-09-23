# -*- coding: utf-8 -*-
"""
Rellena automaticamente informe.docx + justificacion.xlsx con resultados LoadFlow.

Flujo (cualquier alimentador, p.ej. PA217 / IC106 / …):
  1) Copia plantilla limpia (data/output/informe) → doc/
  2) Exige ambos LoadFlow (situacional + proyectado) y meta OCR minima
  3) Genera graficas PNG desde JSON LF (o respeta override manual CYMDIST)
  4) Escribe metricas situacional/proyectado SOLO en celdas de ENTRADA del Excel
     y restaura formulas canónicas de la plantilla original (Informe + C7/C41);
     no fija MD/pérdidas: Excel calcula C185:G186 / VLOOKUP OSM
  5) Extrae tablas evaluando esas formulas y las vuelca 1:1 a las 7 tablas Word
  6) Sustituye textos (conclusión d = nodo de conexión SpotLoad CYMDIST) + imagenes
  7) Render final con Microsoft Word (campos, paginas, PDF)

Imagenes LF obligatorias (auto desde LoadFlow o captura CYMDIST):
  situacional_tension.png | situacional_cargabilidad.png
  proyectado_tension.png  | proyectado_cargabilidad.png

Opcionales: topologia.png | trafo_cargabilidad.png
"""
from __future__ import print_function
import json
import math
import os
import re
import shutil
import zipfile
from datetime import datetime
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from core.common import mkdir, p
from core.feeder_context import load_settings, output_path
from pipeline.assemble_informe import assemble_informe, informe_paths, TEMPLATE_DIR, DOC_DIR
from pipeline.generate_informe_charts import (
    REQUIRED_LF_IMAGES,
    generate_informe_charts,
)
from pipeline.extract_informe_meta_pdf import (
    load_informe_meta,
    meta_is_complete,
)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
ET.register_namespace("w", W_NS)
ET.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")
ET.register_namespace("a", "http://schemas.openxmlformats.org/drawingml/2006/main")
ET.register_namespace("wp", "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing")
ET.register_namespace("pic", "http://schemas.openxmlformats.org/drawingml/2006/picture")
ET.register_namespace("v", "urn:schemas-microsoft-com:vml")

# Plantilla EMAPICA → media a reemplazar si existen capturas nuevas
IMAGE_MAP = {
    "topologia.png": "word/media/image1.png",
    "situacional_tension.png": "word/media/image3.png",
    "situacional_cargabilidad.png": "word/media/image5.png",
    "proyectado_tension.png": "word/media/image7.png",
    "proyectado_cargabilidad.png": "word/media/image9.png",
    "trafo_cargabilidad.png": "word/media/image10.png",
}

LF_OK_STATUS = ("ok", "dry_run")


def _num(val, default=None):
    if val is None or val == "":
        return default
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(" ", "")
    if not s:
        return default
    if s.upper().startswith("ERR"):
        return default
    # Placeholders CYMDIST / macros: $KWLOSS$, $I$, etc.
    if "$" in s:
        return default
    s = s.replace(",", ".")
    try:
        return float(s)
    except Exception:
        return default


def _fmt(val, nd=2, comma=True):
    if val is None:
        return ""
    try:
        x = float(val)
    except Exception:
        return str(val)
    s = ("%." + str(nd) + "f") % x
    if comma:
        s = s.replace(".", ",")
    return s


def _sva(kw, kvar):
    if kw is None or kvar is None:
        return None
    return math.sqrt(kw * kw + kvar * kvar)


def _fp_pct(kw, kva):
    if not kva or kva == 0 or kw is None:
        return None
    return 100.0 * kw / kva


def _amps(kva, vll_kv):
    if not kva or not vll_kv or vll_kv <= 0:
        return None
    return kva / (math.sqrt(3.0) * vll_kv)


def metrics_from_lf(lf_data, settings=None):
    """Normaliza un loadflow_*.json a metricas de informe."""
    lf = lf_data or {}
    topo = dict(lf.get("topo") or {})
    settings = settings or {}
    scen = str(lf.get("scenario") or "").strip().lower()
    # Seguridad: si JSON viejo aún trae KWTOT×3 o sobrelectura leve, corregir aquí
    try:
        from core.cymdist_com import normalize_lf_topo_powers, _parse_com_number
        # Exponer KWTOT crudo situacional para escalar proyectado con el mismo factor
        if scen == "proyectado" and settings.get("_situacional_kw_raw") in (None, ""):
            try:
                from pipeline.assemble_informe import informe_paths
                sit_path = informe_paths(settings).get("loadflow_situacional")
                if sit_path and os.path.isfile(sit_path):
                    with open(sit_path, "r", encoding="utf-8") as f:
                        sit_raw_topo = (json.load(f) or {}).get("topo") or {}
                    settings = dict(settings)
                    settings["_situacional_kw_raw"] = _parse_com_number(
                        sit_raw_topo.get("KWTOT_raw") or sit_raw_topo.get("KWTOT")
                    )
            except Exception:
                pass
        topo = normalize_lf_topo_powers(topo, settings, scenario=scen)
    except Exception:
        pass
    kw = _num(topo.get("KWTOT"))
    kvar = _num(topo.get("KVARTOT"))
    kw_loss = _num(topo.get("KWLOSS"))
    kvar_loss = _num(topo.get("KVARLOSS"))
    vll = _num(topo.get("VLL")) or _num(settings.get("voltage_ll_kv")) or 22.9
    vln = _num(topo.get("VLN"))
    if vln is None and vll:
        vln = vll / math.sqrt(3.0)
    vpu = _num(topo.get("Vpu"))
    vpua = _num(topo.get("VpuA"), vpu)
    vpub = _num(topo.get("VpuB"), vpu)
    vpuc = _num(topo.get("VpuC"), vpu)
    # Corregir Vpu absurdo (>1.2) usando VLN / (VLL/sqrt3)
    nom_vln = (vll / math.sqrt(3.0)) if vll else None
    if nom_vln and vln and (vpu is None or vpu > 1.2 or vpu < 0.5):
        vpu = vln / nom_vln
        vpua = vpua if (vpua and 0.5 <= vpua <= 1.2) else vpu
        vpub = vpub if (vpub and 0.5 <= vpub <= 1.2) else vpu
        vpuc = vpuc if (vpuc and 0.5 <= vpuc <= 1.2) else vpu
    kva = _sva(kw, kvar) if (kw is not None and kvar is not None) else None
    i_a = _num(topo.get("I")) or _amps(kva, vll)
    # Pérdidas: None si placeholder/ausente (no fingir 0.0)
    kva_loss = _sva(kw_loss, kvar_loss) if (kw_loss is not None and kvar_loss is not None) else None
    if kva_loss is None and kw_loss is not None and kvar_loss is None:
        kva_loss = abs(kw_loss)
    elif kva_loss is None and kvar_loss is not None and kw_loss is None:
        kva_loss = abs(kvar_loss)
    fp_loss = None
    if kw_loss is not None and kva_loss:
        fp_loss = _fp_pct(kw_loss, kva_loss)
    return {
        "status": lf.get("status"),
        "scenario": lf.get("scenario"),
        "feeder_id": lf.get("feeder_id") or settings.get("feeder_id"),
        "network_id": lf.get("network_id") or settings.get("network_id"),
        "engine": lf.get("engine"),
        "kw": kw,
        "kvar": kvar,
        "kva": kva,
        "fp_pct": _fp_pct(kw, kva),
        "kw_loss": kw_loss,
        "kvar_loss": kvar_loss,
        "kva_loss": kva_loss,
        "fp_loss_pct": fp_loss,
        "vll": vll,
        "vln": vln,
        "vpu": vpu,
        "vpu_a": vpua,
        "vpu_b": vpub,
        "vpu_c": vpuc,
        "v_pct_a": (vpua * 100.0) if vpua is not None else None,
        "v_pct_b": (vpub * 100.0) if vpub is not None else None,
        "v_pct_c": (vpuc * 100.0) if vpuc is not None else None,
        "i_a": i_a,
        "source_node": lf.get("source_node"),
        "power_scale_reason": topo.get("power_scale_reason"),
        "raw_topo": topo,
    }


def _read_json(path):
    if not path or not os.path.isfile(path):
        return None
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def _load_scenarios(settings):
    paths = informe_paths(settings)
    sit = _read_json(paths["loadflow_situacional"])
    proy = _read_json(paths["loadflow_proyectado"])
    gen = _read_json(paths["loadflow_result"])
    # Fallback: un solo flujo general alimenta situacional
    if sit is None and gen and gen.get("status") in ("ok", "dry_run"):
        sit = dict(gen)
        sit["scenario"] = "situacional"
        sit["_fallback_from"] = "loadflow_result.json"
    if proy is None and gen and str(gen.get("scenario") or "").lower() == "proyectado":
        proy = dict(gen)
    return {
        "situacional": metrics_from_lf(sit, settings) if sit else None,
        "proyectado": metrics_from_lf(proy, settings) if proy else None,
        "raw_situacional": sit,
        "raw_proyectado": proy,
        "paths": paths,
    }


def _lf_status_ok(raw):
    if not raw:
        return False
    return str(raw.get("status") or "").lower() in LF_OK_STATUS


def _meta_cliente(settings):
    """Cabecera del informe: prioriza OCR (informe_meta.json), luego spot loads."""
    feeder = settings.get("feeder_id") or settings.get("active_feeder") or ""
    meta = {
        "cliente": "",
        "ubicacion": settings.get("region") or "Electro Dunas",
        "solicitud": "Factibilidad y Punto de Diseño",
        "potencia_txt": "",
        "potencia_kw": None,
        "alimentador": feeder,
        "set": settings.get("substation") or "",
        "tension_kv": _num(settings.get("voltage_ll_kv"), 22.9),
        "transformador": settings.get("transformer") or "",
        "expediente": settings.get("expediente") or (("RECYM-%s" % feeder) if feeder else "RECYM"),
        "meta_source": "defaults",
    }
    ocr = load_informe_meta(settings) or {}
    if ocr:
        meta["meta_source"] = ocr.get("source") or "pdf"
        for key in (
            "cliente", "ubicacion", "solicitud", "potencia_txt", "alimentador",
            "set", "transformador", "expediente", "proyecto",
        ):
            val = ocr.get(key)
            if val is not None and str(val).strip() != "":
                meta[key] = val
        if ocr.get("potencia_kw") is not None:
            meta["potencia_kw"] = _num(ocr.get("potencia_kw"))
        if ocr.get("tension_kv") is not None:
            meta["tension_kv"] = _num(ocr.get("tension_kv"), meta["tension_kv"])
        if ocr.get("source_pdf"):
            meta["source_pdf"] = ocr.get("source_pdf")

    # Spot loads: complementan potencia / cliente solo si OCR no los trajo
    report = output_path(settings, "loads", "new_spot_loads_report.csv")
    if os.path.isfile(report):
        try:
            import csv
            rows = []
            with open(report, "r", encoding="utf-8") as f:
                for r in csv.DictReader(f):
                    if str(r.get("dry_run") or "").lower() in ("1", "true", "yes"):
                        continue
                    if str(r.get("Estado") or "").upper() in ("DRY_RUN", "ERROR"):
                        continue
                    pk = _num(r.get("P_kW"))
                    if pk is not None:
                        rows.append((pk, r))
            if rows:
                total = sum(x[0] for x in rows)
                if meta.get("potencia_kw") is None:
                    meta["potencia_kw"] = total
                    meta["potencia_txt"] = "%sKW" % _fmt(total, 0, comma=False)
                if not (meta.get("cliente") or "").strip():
                    meta["cliente"] = "Nueva SpotLoad x%d — %s" % (len(rows), feeder)
                    last = rows[-1][1]
                    if last.get("NodeID"):
                        meta["cliente"] = "Carga en nodo %s (%s)" % (last.get("NodeID"), feeder)
        except Exception:
            pass

    if meta.get("potencia_kw") is not None and not (meta.get("potencia_txt") or "").strip():
        meta["potencia_txt"] = "%sKW" % _fmt(meta["potencia_kw"], 0, comma=False)
    if not (meta.get("potencia_txt") or "").strip():
        meta["potencia_txt"] = "—"
    if not (meta.get("alimentador") or "").strip():
        meta["alimentador"] = feeder

    # Punto de diseño = nodo donde se conectó la SpotLoad en CYMDIST
    if not (meta.get("nodo_conexion") or "").strip():
        try:
            from pipeline.generate_location_map import resolve_new_load_node
            loc = resolve_new_load_node(settings)
            nid = (loc or {}).get("NodeID") or ""
            if nid:
                meta["nodo_conexion"] = str(nid).strip()
                meta["nodo_conexion_source"] = (loc or {}).get("source") or ""
        except Exception:
            pass

    return _complete_meta_fields(meta, settings)


def _complete_meta_fields(meta, settings=None):
    """Completa SET/trafo/ubicacion/etc. sin romper campos OCR ya llenos."""
    import re
    settings = settings or {}
    m = dict(meta or {})
    feeder = (m.get("alimentador") or settings.get("feeder_id") or "").strip().upper()
    if feeder:
        m["alimentador"] = feeder
    if not (m.get("solicitud") or "").strip():
        m["solicitud"] = "Factibilidad y Punto de Diseño"
    if m.get("tension_kv") is None:
        m["tension_kv"] = _num(settings.get("voltage_ll_kv"), 22.9)
    if not (m.get("set") or "").strip():
        prefix = re.match(r"^([A-Z]+)", feeder or "")
        geo = {
            "PA": "SET Paracas", "PI": "SET Pisco", "IC": "SET Ica",
            "CH": "SET Chincha", "NA": "SET Nazca",
        }
        m["set"] = geo.get((prefix.group(1) if prefix else ""), "SET %s" % (feeder or "MT"))
    if not (m.get("transformador") or "").strip():
        vll = _num(m.get("tension_kv"), 22.9)
        m["transformador"] = "Transformador %s · %.1f kV (inventario CYMDIST)" % (
            m.get("set") or "SET", vll,
        )
    if not (m.get("expediente") or "").strip():
        m["expediente"] = ("RECYM-%s" % feeder) if feeder else "RECYM"
    if not (m.get("ubicacion") or "").strip():
        m["ubicacion"] = settings.get("region") or "Área de concesión Electro Dunas"
    if m.get("potencia_kw") is not None and not (m.get("potencia_txt") or "").strip():
        m["potencia_txt"] = "%dKW" % int(round(float(m["potencia_kw"])))
    if not (m.get("sistema_electrico") or "").strip():
        m["sistema_electrico"] = "Electro Dunas"
    if not (m.get("codigo_se") or "").strip():
        m["codigo_se"] = feeder or "SE-MT"
    return m


def _images_dir(settings):
    out_base = settings.get("output_dir") or os.path.join(
        "data", "output", "feeders", str(settings.get("feeder_id") or "feeder")
    )
    if not os.path.isabs(out_base):
        return p(*(out_base.replace("\\", "/").split("/") + ["informe_images"]))
    return os.path.join(out_base, "informe_images")


def _check_delivery_gates(scenarios, meta, img_dir, charts_res=None, settings=None):
    """Valida requisitos de entrega rigurosa. Retorna (ready, missing)."""
    missing = []
    raw_sit = scenarios.get("raw_situacional")
    raw_proy = scenarios.get("raw_proyectado")
    if not scenarios.get("situacional") or not _lf_status_ok(raw_sit):
        missing.append("loadflow_situacional")
    if not scenarios.get("proyectado") or not _lf_status_ok(raw_proy):
        missing.append("loadflow_proyectado")
    # Meta debe existir en informe_meta.json (PDF OCR / UI §5), no solo spot-load
    ocr_meta = load_informe_meta(settings) if settings is not None else None
    if not meta_is_complete(ocr_meta):
        missing.append("informe_meta_ocr (cliente + potencia_kw via PDF §5)")
    for fname in REQUIRED_LF_IMAGES:
        if not os.path.isfile(os.path.join(img_dir, fname)):
            missing.append(fname)
    if charts_res and charts_res.get("errors"):
        for err in charts_res["errors"]:
            missing.append("chart_error:%s" % err)
    return (len(missing) == 0), missing


# Celdas con fórmula en «Resultados de escenarios» (NUNCA sobrescribir con valor).
RESULTADOS_FORMULA_CELLS = frozenset({"C7", "C41"})

# Fórmulas canónicas de la plantilla (data/output/informe/justificacion.xlsx).
# Se restauran en cada fill si faltan o si alguien las reemplazó por un número fijo.
# Solo se escriben ENTRADAS LF en Resultados; Informe calcula por estas fórmulas.
RESULTADOS_CANONICAL_FORMULAS = {
    "C7": "=0.7*C6*C6+0.3*C6",       # Fperdida situacional
    "C41": "=0.7*C40*C40+0.3*C40",   # Fperdida proyectado
}

INFORME_CANONICAL_FORMULAS = {
    # Antecedentes (cabecera → texto)
    "B14": '=CONCATENATE("El ",H7,", la Unidad de Proyectos y Obras distribución solicita evaluación de ",C6, " en media tensión, predio ubicado en el ",C5,".")',
    "B15": '=CONCATENATE("La solicitud esta registrado con ",C10)',
    "B16": '=CONCATENATE("El interesado solicita ",C6," ",H6,".")',
    "B17": '=CONCATENATE("El interesado se encuentra próximo al alimentador ",C8,".")',
    "B18": '=CONCATENATE("El transformador de la SET ",H8," tiene las siguientes características ",H9,".")',
    # Situacional: producción / pérdidas / caída (cuadros Word 0–1)
    "E32": "='Resultados de escenarios'!O14",
    "F32": "='Resultados de escenarios'!P14",
    "G32": "='Resultados de escenarios'!Q14",
    "H32": "='Resultados de escenarios'!R14",
    "E33": "='Resultados de escenarios'!O28",
    "F33": "='Resultados de escenarios'!P28",
    "G33": "='Resultados de escenarios'!Q28",
    "H33": "='Resultados de escenarios'!R28",
    "C37": "='Resultados de escenarios'!L15",
    "C38": "='Resultados de escenarios'!L16",
    "C39": "='Resultados de escenarios'!L17",
    # Proyectado: producción / pérdidas / caída (cuadros Word 2, 4)
    "E107": "='Resultados de escenarios'!O48",
    "F107": "='Resultados de escenarios'!P48",
    "G107": "='Resultados de escenarios'!Q48",
    "H107": "='Resultados de escenarios'!R48",
    "E108": "='Resultados de escenarios'!O62",
    "F108": "='Resultados de escenarios'!P62",
    "G108": "='Resultados de escenarios'!Q62",
    "H108": "='Resultados de escenarios'!R62",
    "C112": "='Resultados de escenarios'!L49",
    "C113": "='Resultados de escenarios'!L50",
    "C114": "='Resultados de escenarios'!L51",
    # OSM / pérdidas reconocidas → bloque MD
    "C179": "=C8",
    "D179": "=VLOOKUP(C8,'Clientes por radial_V2'!C6:E98,3,0)",
    "E179": "=VLOOKUP(C8,'Clientes por radial_V2'!C6:D98,2,0)",
    "F179": "=VLOOKUP(E179,Perdidas_ELDU!D6:H32,5,0)",
    "G179": "=VLOOKUP(E179,Perdidas_ELDU!D7:H32,4,0)",
    # Tabla MD / pérdidas (cuadro Word perdidas_md) — no fijar valores aquí
    "C185": "='Resultados de escenarios'!H5",
    "D185": "=C185*24*365*'Resultados de escenarios'!$C$6",
    "E185": "='Resultados de escenarios'!D16*1000",
    "F185": "=E185/D185",
    "G185": "=-(($G$179/100)*D185-E185)*$G$183",
    "C186": "='Resultados de escenarios'!H39",
    "D186": "=C186*24*365*'Resultados de escenarios'!$C$40",
    "E186": "='Resultados de escenarios'!D50*1000",
    "F186": "=E186/D186",
    "G186": "=-(($G$179/100)*D186-E186)*$G$183",
}


# Mapa de celdas de ENTRADA (valores de ejecución LF). El resto de fórmulas
# en hoja Informe referencian estas celdas (O14/H5, L15…, D16, O48/H39…).
EXCEL_SCENARIO_INPUTS = {
    "situacional": {
        "title": "A2",
        "punto": ("C5", "D5", "E5", "F5", "G5", "H5", "I5"),  # Vpu kVLL kVLN I kVA kW kvar
        "fc": "C6",
        "fperdida_formula": "C7",
        "fuente_rows": (12, 14),          # Fuentes / Producción total
        "carga_rows": (15, 16, 20),       # Cargas leída/utilizada/totales
        "vpu_fases": ("L15", "L16", "L17"),
        "loss_rows": (24, 28),            # Pérdidas líneas / totales (resumen O-R)
        "loss_detail_clear": (25, 26, 27),
        "cost_kw": ("C12", "C16"),        # kW pérdidas (entrada)
        "cost_mwh": ("D12", "D16"),       # MW-h/año (derivado de kW×8760×Fperdida)
        "cost_ks": ("E12", "E16"),        # k$/año (derivado)
        "shunt_rows": (21, 22, 23),
        "abnormal_clear": {
            # Limpiar IDs de otro alimentador (plantilla LL202)
            "count": ("J12", "J13", "J14", "J15", "J16", "J17", "J18", "J19", "J20"),
            "worst": ("K12", "K13", "K14", "K15", "K16", "K17", "K18", "K19", "K20"),
            "valor": ("L12", "L13", "L14", "L18", "L19", "L20"),
        },
    },
    "proyectado": {
        "title": "B34",
        "punto": ("C39", "D39", "E39", "F39", "G39", "H39", "I39"),
        "fc": "C40",
        "fperdida_formula": "C41",
        "fuente_rows": (46, 48),
        "carga_rows": (49, 50, 54),
        "vpu_fases": ("L49", "L50", "L51"),
        "loss_rows": (58, 62),
        "loss_detail_clear": (59, 60, 61),
        "cost_kw": ("C46", "C50"),
        "cost_mwh": ("D46", "D50"),
        "cost_ks": ("E46", "E50"),
        "shunt_rows": (55, 56, 57),
        "abnormal_clear": {
            "count": ("J46", "J47", "J48", "J49", "J50", "J51", "J52", "J53", "J54"),
            "worst": ("K46", "K47", "K48", "K49", "K50", "K51", "K52", "K53", "K54"),
            "valor": ("L46", "L47", "L48", "L52", "L53", "L54"),
        },
    },
}


def _cell_is_formula(ws, coord):
    v = ws[coord].value
    return isinstance(v, str) and v.startswith("=")


def _safe_set(ws, coord, value, protected=None):
    """Escribe valor solo si la celda no es fórmula (ni está en protected)."""
    protected = protected or RESULTADOS_FORMULA_CELLS
    if coord in protected or _cell_is_formula(ws, coord):
        return False
    ws[coord].value = value
    return True


def _round_or_none(val, nd=2):
    if val is None:
        return None
    try:
        return round(float(val), nd)
    except Exception:
        return None


def _fperdida_from_fc(fc):
    """Misma expresión que C7/C41: 0.7·FC² + 0.3·FC."""
    try:
        x = float(fc)
    except Exception:
        x = 0.74
    return 0.7 * x * x + 0.3 * x


def _loss_mwh_year(kw_loss, fc):
    if kw_loss is None:
        return None
    return float(kw_loss) * 8760.0 * _fperdida_from_fc(fc) / 1000.0


def _set_or_pq(ws, row, kw, kvar, kva, fp):
    _safe_set(ws, "O%d" % row, _round_or_none(kw, 2))
    _safe_set(ws, "P%d" % row, _round_or_none(kvar, 2))
    _safe_set(ws, "Q%d" % row, _round_or_none(kva, 2))
    _safe_set(ws, "R%d" % row, _round_or_none(fp, 2))


def _clear_or_pq(ws, row):
    _set_or_pq(ws, row, None, None, None, None)


def _write_escenario_block(ws, m, mode):
    """
    Rellena celdas de ENTRADA del bloque situacional/proyectado.
    No toca fórmulas (C7/C41). Limpia residuos de la plantilla (otro feeder)
    cuando no hay dato de ejecución.
    """
    if not m:
        return []
    cfg = EXCEL_SCENARIO_INPUTS.get(mode)
    if not cfg:
        return []
    notes = []
    feeder = m.get("feeder_id") or ""
    has_loss = m.get("kw_loss") is not None or m.get("kvar_loss") is not None
    fc_coord = cfg["fc"]
    fc_val = ws[fc_coord].value
    if fc_val is None or (isinstance(fc_val, str) and not fc_val.strip()):
        _safe_set(ws, fc_coord, 0.74)
        fc_val = 0.74

    title = (
        "ESCENARIO ACTUAL DEL ALIMENTADOR %s" % feeder
        if mode == "situacional"
        else "ESCENARIO PROYECTADO DEL ALIMENTADOR %s" % feeder
    )
    _safe_set(ws, cfg["title"], title)

    # Punto de diseño / cabecera eléctrica
    punto_vals = (
        _round_or_none(m.get("vpu"), 4),
        _round_or_none(m.get("vll"), 2),
        _round_or_none(m.get("vln"), 2),
        _round_or_none(m.get("i_a"), 1),
        _round_or_none(m.get("kva"), 2),
        _round_or_none(m.get("kw"), 2),
        _round_or_none(m.get("kvar"), 2),
    )
    for coord, val in zip(cfg["punto"], punto_vals):
        _safe_set(ws, coord, val)

    # Fuentes / producción = demanda en cabecera
    for row in cfg["fuente_rows"]:
        _set_or_pq(ws, row, m.get("kw"), m.get("kvar"), m.get("kva"), m.get("fp_pct"))

    # Cargas: fuente − pérdidas si hay dato; si no, ≈ fuente (no dejar plantilla ajena)
    if m.get("kw") is not None:
        if has_loss:
            load_kw = m["kw"] - (m.get("kw_loss") or 0)
            load_kvar = (
                (m["kvar"] - (m.get("kvar_loss") or 0))
                if m.get("kvar") is not None
                else None
            )
        else:
            load_kw = m["kw"]
            load_kvar = m.get("kvar")
        load_kva = _sva(load_kw, load_kvar) if load_kvar is not None else m.get("kva")
        load_fp = _fp_pct(load_kw, load_kva)
        for row in cfg["carga_rows"]:
            _set_or_pq(ws, row, load_kw, load_kvar, load_kva, load_fp)

    # Vpu por fase (entrada → Informe C37–C39 / C112–C114)
    for coord, key in zip(cfg["vpu_fases"], ("vpu_a", "vpu_b", "vpu_c")):
        _safe_set(ws, coord, _round_or_none(m.get(key), 4))

    # Pérdidas resumen + costo anual
    if has_loss:
        for row in cfg["loss_rows"]:
            _set_or_pq(
                ws, row,
                m.get("kw_loss"), m.get("kvar_loss"),
                m.get("kva_loss"), m.get("fp_loss_pct"),
            )
        for row in cfg["loss_detail_clear"]:
            _clear_or_pq(ws, row)
        mwh = _loss_mwh_year(m.get("kw_loss"), fc_val)
        for coord in cfg["cost_kw"]:
            _safe_set(ws, coord, _round_or_none(m.get("kw_loss"), 2))
        for coord in cfg["cost_mwh"]:
            _safe_set(ws, coord, _round_or_none(mwh, 2))
        for coord in cfg["cost_ks"]:
            # Plantilla histórica: k$/año ≈ 0.1 × MW-h/año
            _safe_set(ws, coord, _round_or_none((mwh or 0) * 0.1, 2) if mwh is not None else None)
    else:
        # Sin KWLOSS de ejecución: borrar residuos de plantilla (otro estudio)
        for row in list(cfg["loss_rows"]) + list(cfg["loss_detail_clear"]):
            _clear_or_pq(ws, row)
        for coord in list(cfg["cost_kw"]) + list(cfg["cost_mwh"]) + list(cfg["cost_ks"]):
            _safe_set(ws, coord, None)
        notes.append("%s: perdidas no disponibles (placeholders COM)" % mode)

    # Capacitancia shunt: sin dato LF → limpiar (evitar números de LL202)
    for row in cfg["shunt_rows"]:
        _clear_or_pq(ws, row)

    # Condiciones anormales de otro alimentador
    abn = cfg.get("abnormal_clear") or {}
    for coord in abn.get("count") or ():
        _safe_set(ws, coord, 0)
    for coord in abn.get("worst") or ():
        _safe_set(ws, coord, None)
    for coord in abn.get("valor") or ():
        _safe_set(ws, coord, None)

    notes.append("%s: kW=%s kvar=%s loss=%s" % (
        mode, m.get("kw"), m.get("kvar"),
        m.get("kw_loss") if has_loss else "n/d",
    ))
    return notes


def _formula_norm(s):
    return (s or "").replace(" ", "").upper()


def _ensure_canonical_formulas(ws, formulas, protected_inputs=None):
    """
    Restaura fórmulas de la plantilla original si la celda:
      - está vacía, o
      - tiene un valor fijo (ya no es fórmula), o
      - es fórmula distinta a la canónica.
    No toca celdas listadas en protected_inputs (entradas LF).
    """
    protected_inputs = protected_inputs or frozenset()
    n = 0
    for coord, formula in formulas.items():
        if coord in protected_inputs:
            continue
        cur = ws[coord].value
        if isinstance(cur, str) and cur.startswith("="):
            if _formula_norm(cur) == _formula_norm(formula):
                continue
            ws[coord] = formula
            n += 1
        else:
            # Vacío o número fijo → volver a fórmula de plantilla
            ws[coord] = formula
            n += 1
    return n


def _repair_informe_formulas(wi):
    """Compat: asegura fórmulas canónicas Informe (plantilla original)."""
    return _ensure_canonical_formulas(wi, INFORME_CANONICAL_FORMULAS)


def fill_excel(xlsx_path, scenarios, meta):
    """
    Actualiza justificacion.xlsx para CUALQUIER alimentador:

      1) Copia ya vino de plantilla (assemble); aquí solo se rellenan ENTRADAS LF
      2) Se restauran fórmulas canónicas Informe + Resultados (C7/C41) si faltan
         o si quedaron valores fijos — mismas expresiones que el Excel original
      3) Etiquetas / cabecera con meta del alimentador en análisis
      4) excel_tables evalúa la cadena de fórmulas (MD/pérdidas) para el Word

    Nunca escribe números en C185:G186 ni en VLOOKUPs OSM: eso lo calcula Excel.
    """
    wb = load_workbook(xlsx_path)
    notes = []
    if "Resultados de escenarios" in wb.sheetnames:
        ws = wb["Resultados de escenarios"]
        n_res = _ensure_canonical_formulas(ws, RESULTADOS_CANONICAL_FORMULAS)
        if n_res:
            notes.append("Resultados: %d formulas Fperdida restauradas (C7/C41)" % n_res)
        notes += _write_escenario_block(ws, scenarios.get("situacional"), "situacional")
        notes += _write_escenario_block(ws, scenarios.get("proyectado"), "proyectado")
    if "Informe" in wb.sheetnames:
        wi = wb["Informe"]
        # Primero fórmulas (por si una corrida previa las convirtió en valores)
        nfix = _ensure_canonical_formulas(wi, INFORME_CANONICAL_FORMULAS)
        if nfix:
            notes.append("Informe: %d formulas canónicas restauradas (plantilla)" % nfix)
        # Cabecera: solo celdas de entrada (B14–B18 / OSM / MD son fórmulas)
        for coord, val in (
            ("C4", meta.get("cliente") or ""),
            ("C5", meta.get("ubicacion") or ""),
            ("C6", meta.get("solicitud") or "Factibilidad y Punto de Diseño"),
            ("H6", meta.get("potencia_txt") or ""),
            ("C8", meta.get("alimentador") or ""),
            ("H8", meta.get("set") or ""),
            ("H9", meta.get("transformador") or ""),
            ("C10", meta.get("expediente") or ""),
        ):
            if not _cell_is_formula(wi, coord):
                wi[coord] = val
        if meta.get("potencia_kw") is not None and not _cell_is_formula(wi, "C7"):
            wi["C7"] = meta["potencia_kw"]
        if not _cell_is_formula(wi, "H7"):
            wi["H7"] = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
        if meta.get("tension_kv") is not None and not _cell_is_formula(wi, "C9"):
            wi["C9"] = meta["tension_kv"]
        feeder = meta.get("alimentador") or ""
        # Etiquetas texto (no fórmula) del bloque MD / títulos — cualquier alimentador
        for coord, txt in (
            ("B101", "2.1 RED PROYECTADA DE ELECTRODUNAS %s" % feeder),
            ("B181", "PERDIDAS OBTENIDAS DE ESTUDIOS REALIZADOS AL ALIMENTADOR %s (Incorporando cargas solicitadas)" % feeder),
            ("B185", "ESTADO ACTUAL %s" % feeder),
            ("B186", "ESTADO PROYECTADO %s" % feeder),
        ):
            if not _cell_is_formula(wi, coord):
                wi[coord] = txt
        notes.append("cabecera Informe completa (%s)" % feeder)
    wb.save(xlsx_path)
    tables = extract_excel_tables_for_word(xlsx_path, scenarios, meta)
    return notes, tables


def _vlookup_exact(ws, lookup, first_row, last_row, key_col, return_col):
    """VLOOKUP approximate=FALSE sobre hoja abierta."""
    if ws is None or lookup is None:
        return None
    key = str(lookup).strip()
    for r in range(first_row, last_row + 1):
        cell = ws.cell(r, key_col).value
        if cell is not None and str(cell).strip() == key:
            return ws.cell(r, return_col).value
    return None


def _eval_informe_md_block(wb, mode):
    """
    Evalúa la cadena Informe!C185:G185 (o C186:G186) igual que Excel,
    sin sobrescribir fórmulas. Fuente: Resultados de escenarios + VLOOKUP OSM.
    """
    cfg = EXCEL_SCENARIO_INPUTS[mode]
    ws = wb["Resultados de escenarios"] if "Resultados de escenarios" in wb.sheetnames else None
    wi = wb["Informe"] if "Informe" in wb.sheetnames else None
    if ws is None:
        return {}

    kw = _num(ws[cfg["punto"][5]].value)  # H5 / H39
    fc = _num(ws[cfg["fc"]].value)
    if fc is None:
        fc = 0.74
    # Informe usa D16 / D50 (totales), no D12 / D46
    mwh = _num(ws[cfg["cost_mwh"][-1]].value)
    kw_loss = _num(ws[cfg["cost_kw"][-1]].value)

    md_kwh = None
    if kw is not None:
        md_kwh = float(kw) * 24.0 * 365.0 * float(fc)

    # E185 = D16*1000  (vacío → 0, como Excel)
    loss_kwh = (float(mwh) if mwh is not None else 0.0) * 1000.0

    loss_pct = None  # fracción 0–1 (formato 0.00% en Excel)
    if md_kwh is not None and md_kwh > 0:
        loss_pct = loss_kwh / md_kwh

    # G179 / G183 para costo no reconocido (fórmulas Informe)
    loss_cost = None
    g179 = None
    g183 = _num(wi["G183"].value) if wi is not None else None
    if wi is not None and "Clientes por radial_V2" in wb.sheetnames and "Perdidas_ELDU" in wb.sheetnames:
        feeder = wi["C8"].value
        # E179 = VLOOKUP(C8, Clientes!C6:D98, 2, 0) → col D
        e179 = _vlookup_exact(wb["Clientes por radial_V2"], feeder, 6, 98, 3, 4)
        # G179 = VLOOKUP(E179, Perdidas_ELDU!D7:H32, 4, 0) → col G (índice 4 del rango)
        if e179 is not None:
            g179 = _num(_vlookup_exact(wb["Perdidas_ELDU"], e179, 7, 32, 4, 7))
        if g179 is None:
            # Intentar cached / valor directo si ya no es fórmula
            g179 = _num(wi["G179"].value) if not _cell_is_formula(wi, "G179") else None
    if md_kwh is not None and g179 is not None and g183 is not None:
        # G185 = -(($G$179/100)*D185-E185)*$G$183
        loss_cost = -((float(g179) / 100.0) * float(md_kwh) - float(loss_kwh)) * float(g183)

    return {
        "kw": kw,
        "fc": fc,
        "md_kwh": md_kwh,
        "kw_loss": kw_loss,
        "mwh_loss": mwh,
        "loss_kwh": loss_kwh,
        "loss_pct": loss_pct,
        "loss_cost": loss_cost,
        "g179": g179,
        "from_excel_formulas": True,
    }


def extract_excel_tables_for_word(xlsx_path, scenarios=None, meta=None):
    """
    Extrae de justificacion.xlsx las tablas que corresponden 1:1 a las del Word.
    Entradas LF = celdas Resultados; bloque MD/pérdidas = evaluación de fórmulas Informe.
    """
    wb = load_workbook(xlsx_path, data_only=False)
    ws = wb["Resultados de escenarios"] if "Resultados de escenarios" in wb.sheetnames else None
    meta = meta or {}
    scenarios = scenarios or {}

    def _pq(row):
        if ws is None:
            return {}
        return {
            "kw": _num(ws["O%d" % row].value),
            "kvar": _num(ws["P%d" % row].value),
            "kva": _num(ws["Q%d" % row].value),
            "fp_pct": _num(ws["R%d" % row].value),
        }

    def _v_pct(coord):
        if ws is None:
            return None
        v = _num(ws[coord].value)
        if v is None:
            return None
        # L15… guarda Vp.u.; Word muestra %
        return (v * 100.0) if v <= 1.5 else v

    def _scenario_pack(mode, m_fallback):
        cfg = EXCEL_SCENARIO_INPUTS[mode]
        fuente_row = cfg["fuente_rows"][-1]  # producción total
        loss_row = cfg["loss_rows"][-1]      # pérdidas totales
        punto = cfg["punto"]
        md = _eval_informe_md_block(wb, mode)
        pack = {
            "fuente": _pq(fuente_row),
            "loss": _pq(loss_row),
            "v_pct_a": _v_pct(cfg["vpu_fases"][0]),
            "v_pct_b": _v_pct(cfg["vpu_fases"][1]),
            "v_pct_c": _v_pct(cfg["vpu_fases"][2]),
            "kw": md.get("kw") if md.get("kw") is not None else (
                _num(ws[punto[5]].value) if ws is not None else None
            ),
            "kvar": _num(ws[punto[6]].value) if ws is not None else None,
            "kva": _num(ws[punto[4]].value) if ws is not None else None,
            "vpu": _num(ws[punto[0]].value) if ws is not None else None,
            "vll": _num(ws[punto[1]].value) if ws is not None else None,
            "vln": _num(ws[punto[2]].value) if ws is not None else None,
            "i_a": _num(ws[punto[3]].value) if ws is not None else None,
            "kw_loss": md.get("kw_loss"),
            "mwh_loss": md.get("mwh_loss"),
            "md_kwh": md.get("md_kwh"),
            "loss_kwh": md.get("loss_kwh"),
            "loss_pct": md.get("loss_pct"),
            "loss_cost": md.get("loss_cost"),
            "fp_pct": _pq(fuente_row).get("fp_pct"),
            "fc": md.get("fc") if md.get("fc") is not None else 0.74,
            "from_excel_formulas": md.get("from_excel_formulas"),
        }
        # Fallback a métricas LF si Excel vacío
        fb = m_fallback or {}
        if pack["fuente"].get("kw") is None and fb.get("kw") is not None:
            pack["fuente"] = {
                "kw": fb.get("kw"), "kvar": fb.get("kvar"),
                "kva": fb.get("kva"), "fp_pct": fb.get("fp_pct"),
            }
        if pack["kw"] is None:
            pack["kw"] = fb.get("kw")
        if pack["v_pct_a"] is None and fb.get("v_pct_a") is not None:
            pack["v_pct_a"] = fb.get("v_pct_a")
            pack["v_pct_b"] = fb.get("v_pct_b")
            pack["v_pct_c"] = fb.get("v_pct_c")
        if pack["loss"].get("kw") is None and fb.get("kw_loss") is not None:
            pack["loss"] = {
                "kw": fb.get("kw_loss"), "kvar": fb.get("kvar_loss"),
                "kva": fb.get("kva_loss"), "fp_pct": fb.get("fp_loss_pct"),
            }
            pack["kw_loss"] = fb.get("kw_loss")
            # Re-evaluar kWh/año de pérdidas con FC Excel si LF trae kW loss
            if pack.get("mwh_loss") is None and pack["kw_loss"] is not None:
                pack["mwh_loss"] = _loss_mwh_year(pack["kw_loss"], pack.get("fc"))
                pack["loss_kwh"] = float(pack["mwh_loss"]) * 1000.0
                if pack.get("md_kwh") and pack["md_kwh"] > 0:
                    pack["loss_pct"] = pack["loss_kwh"] / float(pack["md_kwh"])
        return pack

    sit = _scenario_pack("situacional", scenarios.get("situacional"))
    proy = _scenario_pack("proyectado", scenarios.get("proyectado"))
    return {
        "situacional": sit,
        "proyectado": proy,
        "meta": {
            "alimentador": meta.get("alimentador"),
            "sistema_electrico": meta.get("sistema_electrico"),
            "codigo_se": meta.get("codigo_se"),
            "potencia_kw": meta.get("potencia_kw"),
            "tension_kv": meta.get("tension_kv"),
            "nodo_conexion": meta.get("nodo_conexion"),
        },
        "mapping": {
            "situacional_fuente": "Resultados!O14:R14 (+ O28:R28 pérdidas)",
            "situacional_caida": "Resultados!L15:L17 → %",
            "proyectado_fuente": "Resultados!O48:R48 (+ O62:R62 pérdidas)",
            "punto_diseno": "meta.nodo_conexion + meta.potencia_kw + Resultados proyectado V/I",
            "proyectado_caida": "Resultados!L49:L51 → %",
            "osm_alimentador": "Informe!C8 + meta",
            "perdidas_md": "Informe!C185:G186 (=H5/H39, FC, D16/D50, G179, G183)",
        },
    }

def _cell_texts(tc):
    return list(tc.iter("{%s}t" % W_NS))


def _cell_text(tc):
    return "".join((t.text or "") for t in _cell_texts(tc))


def _set_cell_text(tc, value):
    """Escribe texto en una celda: primer w:t recibe el valor; el resto se vacia."""
    if value is None:
        return False
    texts = _cell_texts(tc)
    if not texts:
        # Crear un parrafo/run/t minimo si la celda esta vacia
        p = ET.SubElement(tc, "{%s}p" % W_NS)
        r = ET.SubElement(p, "{%s}r" % W_NS)
        t = ET.SubElement(r, "{%s}t" % W_NS)
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = str(value)
        return True
    texts[0].text = str(value)
    for t in texts[1:]:
        t.text = ""
    return True


def _tbl_rows(tbl):
    rows = []
    for tr in tbl.findall("{%s}tr" % W_NS):
        cells = list(tr.findall("{%s}tc" % W_NS))
        rows.append(cells)
    return rows


def _para_text(p):
    return "".join((t.text or "") for t in p.iter("{%s}t" % W_NS))


def _set_para_text(p, value):
    texts = list(p.iter("{%s}t" % W_NS))
    if not texts:
        return False
    texts[0].text = str(value)
    for t in texts[1:]:
        t.text = ""
    return True


def _replace_in_paragraph(p, old, new):
    """Replace exact substring in a single paragraph (merged runs)."""
    if old is None or new is None or str(old) == str(new):
        return 0
    texts = list(p.iter("{%s}t" % W_NS))
    if not texts:
        return 0
    full = "".join((t.text or "") for t in texts)
    if str(old) not in full:
        return 0
    n = full.count(str(old))
    full = full.replace(str(old), str(new))
    texts[0].text = full
    for t in texts[1:]:
        t.text = ""
    return n


def _merge_w_t(root):
    """Une w:t adyacentes del mismo parrafo para facilitar replace (in-place XML)."""
    for p in root.iter("{%s}p" % W_NS):
        texts = list(p.iter("{%s}t" % W_NS))
        if len(texts) < 2:
            continue
        full = "".join((t.text or "") for t in texts)
        texts[0].text = full
        for t in texts[1:]:
            t.text = ""


def _fmt_md_kw(val):
    """Formato tipo 2,791.00 / 10156.00 para tabla MD."""
    if val is None:
        return None
    try:
        x = float(val)
    except Exception:
        return str(val)
    # miles con coma, decimales con punto (estilo plantilla US)
    int_part = int(abs(x))
    dec = abs(x) - int_part
    s_int = "{:,}".format(int_part)
    s = "%s.%02d" % (s_int, int(round(dec * 100)))
    if x < 0:
        s = "-" + s
    return s


def _punto_diseno_metrics(proy, meta):
    """Metricas del punto de diseno (nueva carga), no de cabecera alimentador."""
    pk = _num(meta.get("potencia_kw"))
    fp = 0.95
    qk = None
    kva = None
    if pk is not None:
        qk = pk * math.tan(math.acos(min(0.999, max(0.01, fp))))
        kva = _sva(pk, qk)
    vll = (proy or {}).get("vll") or _num(meta.get("tension_kv")) or 22.9
    vln = (proy or {}).get("vln")
    if vln is None and vll:
        vln = vll / math.sqrt(3.0)
    # Vp.u. del punto: preferir proy; si absurdo, 1.0 nominal
    vpu = (proy or {}).get("vpu")
    if vpu is None or vpu > 1.2 or vpu < 0.5:
        vpu = 1.0
    i_a = _amps(kva, vll) if kva else None
    return {
        "vpu": vpu,
        "vll": vll,
        "vln": vln,
        "i_a": i_a,
        "kva": kva,
        "kw": pk,
        "kvar": qk,
    }


def _fmt_pq_cell(val, nd_int=0, nd_frac=2):
    if val is None:
        return None
    try:
        x = float(val)
    except Exception:
        return str(val)
    if abs(x - round(x)) < 1e-6:
        return _fmt(x, nd_int, comma=False)
    return _fmt(x, nd_frac, comma=False)


def _fill_source_table(tbl, fuente, loss=None, clear_loss_if_missing=True):
    """Tabla potencia fuente + perdidas (4 filas x 5 cols). Datos desde Excel."""
    done = []
    fuente = fuente or {}
    loss = loss or {}
    rows = _tbl_rows(tbl)
    if len(rows) < 4:
        return done
    pot = rows[2]
    if len(pot) >= 5:
        if fuente.get("kw") is not None and _set_cell_text(pot[1], _fmt_pq_cell(fuente["kw"])):
            done.append("fuente.kw")
        if fuente.get("kvar") is not None and _set_cell_text(pot[2], _fmt_pq_cell(fuente["kvar"])):
            done.append("fuente.kvar")
        if fuente.get("kva") is not None and _set_cell_text(pot[3], _fmt_pq_cell(fuente["kva"])):
            done.append("fuente.kva")
        if fuente.get("fp_pct") is not None and _set_cell_text(pot[4], _fmt(fuente["fp_pct"], 2)):
            done.append("fuente.fp")
    loss_row = rows[3]
    has_loss = loss.get("kw") is not None or loss.get("kvar") is not None
    if has_loss and len(loss_row) >= 5:
        if loss.get("kw") is not None and _set_cell_text(loss_row[1], _fmt(loss["kw"], 2, comma=False)):
            done.append("loss.kw")
        if loss.get("kvar") is not None and _set_cell_text(loss_row[2], _fmt(loss["kvar"], 2, comma=False)):
            done.append("loss.kvar")
        if loss.get("kva") is not None and _set_cell_text(loss_row[3], _fmt(loss["kva"], 2, comma=False)):
            done.append("loss.kva")
        if loss.get("fp_pct") is not None and _set_cell_text(loss_row[4], _fmt(loss["fp_pct"], 2, comma=False)):
            done.append("loss.fp")
    elif clear_loss_if_missing and len(loss_row) >= 5:
        # Evitar residuos de plantilla (otro alimentador)
        for idx, key in ((1, "kw"), (2, "kvar"), (3, "kva"), (4, "fp")):
            if _set_cell_text(loss_row[idx], "—"):
                done.append("loss.%s.cleared" % key)
    return done


def _fill_voltage_drop_table(tbl, pack):
    """Tabla caida A/B/C % (valores ya en % desde extract Excel)."""
    done = []
    pack = pack or {}
    rows = _tbl_rows(tbl)
    mapping = {1: "v_pct_a", 2: "v_pct_b", 3: "v_pct_c"}
    for idx, key in mapping.items():
        if idx >= len(rows):
            break
        cells = rows[idx]
        if len(cells) < 2:
            continue
        val = pack.get(key)
        if val is None:
            continue
        if _set_cell_text(cells[1], _fmt(val, 2, comma=False) + "%"):
            done.append(key)
    return done


def _fill_punto_diseno_table(tbl, pd):
    """Tabla 1 fila datos: Vp.u. kVLL kVLN i KVA kW kvar."""
    done = []
    if not pd:
        return done
    rows = _tbl_rows(tbl)
    if len(rows) < 2:
        return done
    cells = rows[1]
    vals = [
        _fmt(pd.get("vpu"), 2, comma=False) if pd.get("vpu") is not None else None,
        _fmt(pd.get("vll"), 1, comma=False) if pd.get("vll") is not None else None,
        _fmt(pd.get("vln"), 1, comma=False) if pd.get("vln") is not None else None,
        _fmt(pd.get("i_a"), 1, comma=False) if pd.get("i_a") is not None else None,
        _fmt(pd.get("kva"), 0, comma=False) if pd.get("kva") is not None else None,
        _fmt(pd.get("kw"), 0, comma=False) if pd.get("kw") is not None else None,
        _fmt(pd.get("kvar"), 0, comma=False) if pd.get("kvar") is not None else None,
    ]
    keys = ["vpu", "vll", "vln", "i_a", "kva", "kw", "kvar"]
    for i, (key, val) in enumerate(zip(keys, vals)):
        if val is None or i >= len(cells):
            continue
        if _set_cell_text(cells[i], val):
            done.append("pd.%s" % key)
    return done


def _fill_osm_feeder_table(tbl, meta):
    done = []
    feeder = (meta or {}).get("alimentador")
    if not feeder:
        return done
    rows = _tbl_rows(tbl)
    if len(rows) < 2:
        return done
    cells = rows[1]
    vals = [
        None,
        feeder,
        (meta or {}).get("sistema_electrico") or "Electro Dunas",
        (meta or {}).get("codigo_se") or feeder,
    ]
    for i, val in enumerate(vals):
        if val is None or i >= len(cells):
            continue
        if _set_cell_text(cells[i], val):
            done.append("osm.col%d" % i)
    return done


def _fill_perdidas_md_table(tbl, sit, proy, meta):
    """Tabla ESTADO ACTUAL / PROYECTADO: valores = fórmulas Informe!C185:G186."""
    done = []
    rows = _tbl_rows(tbl)
    if len(rows) < 3:
        return done
    feeder = (meta or {}).get("alimentador") or ""
    sit = sit or {}
    proy = proy or {}

    def _fill_md_row(row_cells, label, pack, prefix):
        local = []
        if len(row_cells) < 2:
            return local
        if feeder and _set_cell_text(row_cells[0], label):
            local.append("%s.label" % prefix)
        kw = pack.get("kw")
        if kw is not None and _set_cell_text(row_cells[1], _fmt_md_kw(kw)):
            local.append("%s.kw" % prefix)

        # D185 = C185*24*365*FC  (ya evaluado en extract → md_kwh)
        md_kwh = pack.get("md_kwh")
        fc = pack.get("fc")
        if fc is None:
            fc = 0.74
        if md_kwh is None and kw is not None:
            md_kwh = float(kw) * 24.0 * 365.0 * float(fc)
        if md_kwh is not None and len(row_cells) >= 3:
            if _set_cell_text(row_cells[2], _fmt_md_kw(md_kwh)):
                local.append("%s.kwh" % prefix)

        # E/F/G: mismos resultados que Informe (vacío D16 → 0, no "—")
        if pack.get("from_excel_formulas") or pack.get("loss_kwh") is not None or pack.get("md_kwh") is not None:
            loss_kwh = pack.get("loss_kwh")
            if loss_kwh is None:
                mwh = pack.get("mwh_loss")
                loss_kwh = (float(mwh) if mwh is not None else 0.0) * 1000.0
            if len(row_cells) >= 4 and _set_cell_text(row_cells[3], _fmt_md_kw(loss_kwh)):
                local.append("%s.loss_kwh" % prefix)
            loss_pct = pack.get("loss_pct")
            if loss_pct is None and md_kwh and md_kwh > 0:
                loss_pct = float(loss_kwh) / float(md_kwh)
            if loss_pct is not None and len(row_cells) >= 5:
                # Excel F185 formato 0.00% → valor fracción
                if _set_cell_text(row_cells[4], _fmt(float(loss_pct) * 100.0, 2, comma=False) + "%"):
                    local.append("%s.pct" % prefix)
            loss_cost = pack.get("loss_cost")
            if loss_cost is not None and len(row_cells) >= 6:
                if _set_cell_text(row_cells[5], _fmt_md_kw(loss_cost)):
                    local.append("%s.cost" % prefix)
        elif pack.get("kw_loss") is None and len(row_cells) >= 4:
            if _set_cell_text(row_cells[3], "—"):
                local.append("%s.loss_kwh.cleared" % prefix)
            if len(row_cells) >= 5 and _set_cell_text(row_cells[4], "—"):
                local.append("%s.pct.cleared" % prefix)
            if len(row_cells) >= 6 and _set_cell_text(row_cells[5], "—"):
                local.append("%s.cost.cleared" % prefix)
        return local

    done += _fill_md_row(rows[1], ("ESTADO ACTUAL %s" % feeder).strip(), sit, "md.sit")
    done += _fill_md_row(rows[2], ("ESTADO PROYECTADO %s" % feeder).strip(), proy, "md.proy")
    return done


def _fill_word_tables(root, scenarios, meta, excel_tables=None):
    """
    Rellena los 7 cuadros Word con correspondencia precisa al Excel:
      0 situacional_fuente ← O14:R14 + O28:R28
      1 situacional_caida  ← L15:L17 (% )
      2 proyectado_fuente  ← O48:R48 + O62:R62
      3 punto_diseno       ← meta potencia + V/I proyectado
      4 proyectado_caida   ← L49:L51
      5 osm_alimentador    ← meta alimentador
      6 perdidas_md        ← H5/H39 + D16/D50
    """
    tables = list(root.iter("{%s}tbl" % W_NS))
    xt = excel_tables or {}
    sit_x = xt.get("situacional") or {}
    proy_x = xt.get("proyectado") or {}
    sit = scenarios.get("situacional") or {}
    proy = scenarios.get("proyectado") or {}
    report = {"n_tables": len(tables), "filled": [], "cells": [], "source": "excel" if excel_tables else "metrics"}

    # Packs para tablas fuente/caída: preferir extracción Excel
    sit_fuente = sit_x.get("fuente") or {
        "kw": sit.get("kw"), "kvar": sit.get("kvar"),
        "kva": sit.get("kva"), "fp_pct": sit.get("fp_pct"),
    }
    sit_loss = sit_x.get("loss") or {
        "kw": sit.get("kw_loss"), "kvar": sit.get("kvar_loss"),
        "kva": sit.get("kva_loss"), "fp_pct": sit.get("fp_loss_pct"),
    }
    proy_fuente = proy_x.get("fuente") or {
        "kw": proy.get("kw"), "kvar": proy.get("kvar"),
        "kva": proy.get("kva"), "fp_pct": proy.get("fp_pct"),
    }
    proy_loss = proy_x.get("loss") or {
        "kw": proy.get("kw_loss"), "kvar": proy.get("kvar_loss"),
        "kva": proy.get("kva_loss"), "fp_pct": proy.get("fp_loss_pct"),
    }
    sit_v = {
        "v_pct_a": sit_x.get("v_pct_a", sit.get("v_pct_a")),
        "v_pct_b": sit_x.get("v_pct_b", sit.get("v_pct_b")),
        "v_pct_c": sit_x.get("v_pct_c", sit.get("v_pct_c")),
    }
    proy_v = {
        "v_pct_a": proy_x.get("v_pct_a", proy.get("v_pct_a")),
        "v_pct_b": proy_x.get("v_pct_b", proy.get("v_pct_b")),
        "v_pct_c": proy_x.get("v_pct_c", proy.get("v_pct_c")),
    }
    sit_md = {
        "kw": sit_x.get("kw", sit.get("kw")),
        "kw_loss": sit_x.get("kw_loss", sit.get("kw_loss")),
        "mwh_loss": sit_x.get("mwh_loss"),
        "md_kwh": sit_x.get("md_kwh"),
        "loss_kwh": sit_x.get("loss_kwh"),
        "loss_pct": sit_x.get("loss_pct"),
        "loss_cost": sit_x.get("loss_cost"),
        "fc": sit_x.get("fc", 0.74),
        "from_excel_formulas": sit_x.get("from_excel_formulas"),
    }
    proy_md = {
        "kw": proy_x.get("kw", proy.get("kw")),
        "kw_loss": proy_x.get("kw_loss", proy.get("kw_loss")),
        "mwh_loss": proy_x.get("mwh_loss"),
        "md_kwh": proy_x.get("md_kwh"),
        "loss_kwh": proy_x.get("loss_kwh"),
        "loss_pct": proy_x.get("loss_pct"),
        "loss_cost": proy_x.get("loss_cost"),
        "fc": proy_x.get("fc", 0.74),
        "from_excel_formulas": proy_x.get("from_excel_formulas"),
    }
    # Punto de diseño: potencia solicitada + tensión del escenario proyectado Excel
    pd = _punto_diseno_metrics(proy, meta)
    if proy_x.get("vpu") is not None:
        pd["vpu"] = proy_x["vpu"]
    if proy_x.get("vll") is not None:
        pd["vll"] = proy_x["vll"]
    if proy_x.get("vln") is not None:
        pd["vln"] = proy_x["vln"]

    fillers = [
        ("situacional_fuente", lambda t: _fill_source_table(t, sit_fuente, sit_loss)),
        ("situacional_caida", lambda t: _fill_voltage_drop_table(t, sit_v)),
        ("proyectado_fuente", lambda t: _fill_source_table(t, proy_fuente, proy_loss)),
        ("punto_diseno", lambda t: _fill_punto_diseno_table(t, pd)),
        ("proyectado_caida", lambda t: _fill_voltage_drop_table(t, proy_v)),
        ("osm_alimentador", lambda t: _fill_osm_feeder_table(t, meta)),
        ("perdidas_md", lambda t: _fill_perdidas_md_table(t, sit_md, proy_md, meta)),
    ]
    for i, (name, fn) in enumerate(fillers):
        if i >= len(tables):
            break
        cells = fn(tables[i])
        if cells:
            report["filled"].append(name)
            report["cells"].extend(["%s:%s" % (name, c) for c in cells])
    return report


def _build_header_replacements(meta, scenarios=None):
    """Reemplazos que respetan la estructura de la plantilla base (frases largas)."""
    reps = []
    cliente = (meta or {}).get("cliente") or ""
    alimentador = (meta or {}).get("alimentador") or ""
    ubic = (meta or {}).get("ubicacion") or ""
    pot_kw = _num((meta or {}).get("potencia_kw"))
    pot_txt = (meta or {}).get("potencia_txt") or ""
    expediente = (meta or {}).get("expediente") or ""
    set_name = (meta or {}).get("set") or ""
    trafo = (meta or {}).get("transformador") or ""
    proyecto = (meta or {}).get("proyecto") or ""
    nodo = (meta or {}).get("nodo_conexion") or ""
    if pot_kw is not None and not pot_txt:
        pot_txt = "%sKW" % _fmt(pot_kw, 0, comma=False)
    pot_display = ("%s kW" % _fmt(pot_kw, 0, comma=False)) if pot_kw is not None else (pot_txt or "")

    # Fecha cabecera
    meses = (
        "enero", "febrero", "marzo", "abril", "mayo", "junio",
        "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
    )
    now = datetime.now()
    fecha_txt = "Ica, %d de %s de %d" % (now.day, meses[now.month - 1], now.year)
    reps.append(("Ica, 14 de agosto de 2026", fecha_txt))

    # Frases LARGAS primero (antes de IC106 / 430 kW / EMAPICA sueltos)
    if cliente and pot_display:
        label = "Carga-%s-%s" % ((proyecto or cliente)[:40], pot_display)
        if nodo:
            label = "%s- Nodo %s" % (label, nodo)
        old_carga = "Carga-Bomba de Agua Residual-juan Santa-430 kW- Cod IGEA-580224223"
        # La plantilla a veces duplica el rótulo en el mismo párrafo
        reps.append((old_carga + old_carga, label))
        reps.append((old_carga, label))
    if proyecto:
        reps.append(("Bomba de agua residual - Juan Santa", proyecto))
        reps.append(("Bomba de Agua Residual-juan Santa", proyecto))
    if cliente and alimentador and pot_display:
        concl1 = (
            "La solicitud de factibilidad y asignación de punto de diseño presentada por el predio %s, "
            "para una demanda de %s, se encuentra dentro del área de concesión de la empresa distribuidora."
        ) % (cliente, pot_display)
        concl2 = (
            "Asimismo, se evaluó el estado actual (sin proyecto) y proyectado (con proyecto) de cargabilidad "
            "y calidad de producto del alimentador %s y de la %s, determinándose que cuentan con condiciones "
            "operativas favorables para la incorporación de la nueva carga solicitada. Los resultados del "
            "flujo de carga evidencian capacidad suficiente para atender la demanda requerida y mantienen "
            "niveles de tensión dentro de los rangos establecidos por la normativa vigente."
        ) % (alimentador, set_name or "SET asociada")
        if nodo:
            punto_txt = (
                "en el nodo %s del alimentador %s (punto de conexión de la nueva carga en CYMDIST)"
                % (nodo, alimentador)
            )
        else:
            punto_txt = "en el alimentador %s%s" % (
                alimentador,
                (" — proyecto «%s»" % proyecto) if proyecto else "",
            )
        concl3 = (
            "En consecuencia, para la atención de la demanda requerida por el predio %s, se recomienda "
            "asignar el punto de diseño %s. Asimismo, se verifica que tanto el "
            "alimentador como el transformador de potencia de la %s cuentan con capacidad disponible "
            "para el suministro solicitado, manteniéndose los niveles de tensión dentro de los límites "
            "permisibles y sin afectar la calidad de producto ni las condiciones operativas de la red eléctrica."
        ) % (
            cliente,
            punto_txt,
            set_name or "SET",
        )
        # Plantilla base trae un párrafo de rechazo (IC107): puente breve (no duplicar concl2)
        reps.append((
            "De acuerdo con los análisis de cargabilidad y calidad de producto realizados, no es posible asignar el punto de diseño en el alimentador IC107, debido a que este presenta niveles de cargabilidad cercanos a su capacidad máxima de operación y condiciones de caída de tensión próximas a los límites permisibles establecidos por la normativa vigente",
            (
                "De acuerdo con los análisis de cargabilidad y calidad de producto realizados sobre el alimentador %s "
                "(estados situacional y proyectado), se verifica que la red mantiene niveles de tensión y "
                "cargabilidad dentro de los parámetros operativos permitidos para la demanda solicitada."
            ) % alimentador,
        ))
        reps.append((
            "La solicitud de factibilidad y asignación de punto de diseño presentada por el predio MUNICIPAL DE AGUA POTABLE Y ALCANTARILLADO DE ICA S.A. - EMAPICA, para una demanda de 430 kW, se encuentra dentro del área de concesión de la empresa distribuidora.",
            concl1,
        ))
        reps.append((
            "para una demanda de 430 kW, se encuentra dentro del área de concesión de la empresa distribuidora.",
            "para una demanda de %s, se encuentra dentro del área de concesión de la empresa distribuidora." % pot_display,
        ))
        reps.append((
            "Asimismo, se evaluó el estado actual de cargabilidad y calidad de producto del alimentador IC106 y del transformador de potencia de la SET Ica, determinándose que ambos cuentan con condiciones operativas favorables para la incorporación de la nueva carga solicitada. Los resultados obtenidos evidencian que el alimentador IC106 dispone de capacidad suficiente para atender la demanda requerida y mantiene niveles de tensión dentro de los rangos establecidos por la normativa vigente. De igual manera, el transformador de potencia de la SET Ica presenta disponibilidad de capacidad, garantizando una operación segura y confiable del sistema.",
            concl2,
        ))
        reps.append((
            "En consecuencia, para la atención de la demanda requerida por el predio MUNICIPAL DE AGUA POTABLE Y ALCANTARILLADO DE ICA S.A. - EMAPICA, se recomienda asignar el punto de diseño en la estructura de media tensión NMT N.° COD IGEA: 580224223. Asimismo, se verifica que tanto el alimentador IC106 como el transformador de potencia de la SET Ica cuentan con capacidad disponible para el suministro solicitado, manteniéndose los niveles de tensión dentro de los límites permisibles y sin afectar la calidad de producto ni las condiciones operativas de la red eléctrica",
            concl3,
        ))

    if alimentador:
        reps.append(("IC106", alimentador))
        reps.append(("IC107", alimentador))
    if cliente:
        reps.append((
            "MUNICIPAL DE AGUA POTABLE Y ALCANTARILLADO DE ICA S.A. - EMAPICA",
            cliente,
        ))
        reps.append((
            "PREDIO MUNICIPAL DE AGUA POTABLE Y ALCANTARILLADO DE ICA S.A. - EMAPICA",
            cliente,
        ))
        reps.append(("EMAPICA", cliente[:80]))
    if pot_display:
        reps.append(("430 KW", pot_display))
        reps.append(("430 kW", pot_display))
        if pot_txt:
            reps.append(("430KW", pot_txt))
        reps.append((
            "El interesado solicita Factibilidad y Punto de Diseño 430 KW.",
            "El interesado solicita Factibilidad y Punto de Diseño %s." % pot_display,
        ))
        reps.append((
            "El interesado solicita Factibilidad y Punto de Diseño 430 kW.",
            "El interesado solicita Factibilidad y Punto de Diseño %s." % pot_display,
        ))
    if expediente:
        reps.append(("EXP-2026-000434", expediente))
    if set_name:
        reps.append(("SET ICA", set_name))
        reps.append(("SET Ica", set_name))
    if trafo:
        reps.append((
            "El transformador de la SET ICA tiene las siguientes características 50/30/30 MVA.",
            "El transformador de la %s tiene las siguientes características: %s." % (set_name or "SET", trafo),
        ))
        reps.append(("50/30/30 MVA", trafo))
    if alimentador:
        reps.append((
            "El interesado se encuentra próximo al alimentador IC106.",
            "El interesado se encuentra próximo al alimentador %s." % alimentador,
        ))
        reps.append((
            "Se ha verificado que el interesado se encuentra dentro de área de la zona de influencia del alimentador IC106",
            "Se ha verificado que el interesado se encuentra dentro del área de influencia del alimentador %s%s." % (
                alimentador,
                (" (%s)" % ubic) if ubic else "",
            ),
        ))
    if ubic:
        reps.append((
            "distrito de Parcona, provincia ICA del departamento de ICA",
            ubic,
        ))

    return reps


def _fill_word_headers(root, meta, scenarios=None):
    """Actualiza ASUNTO / ANTECEDENTES / CONCLUSIONES con meta (estructura plantilla)."""
    _merge_w_t(root)
    reps = _build_header_replacements(meta, scenarios)
    done = []
    for old, new in reps:
        count = 0
        for p in root.iter("{%s}p" % W_NS):
            count += _replace_in_paragraph(p, old, new)
        if count:
            done.append({"old": old[:80], "new": str(new)[:80], "count": count})
    cliente = (meta or {}).get("cliente") or ""
    if cliente:
        for p in root.iter("{%s}p" % W_NS):
            txt = _para_text(p)
            if "ASUNTO" in txt.upper() and "FACTIBILIDAD" in txt.upper():
                new_txt = "ASUNTO: FACTIBILIDAD AMPLIACION DE CARGA, %s." % cliente
                if _set_para_text(p, new_txt):
                    done.append({"old": "ASUNTO…", "new": new_txt, "count": 1})
                break
    # Antecedentes ubicacion (evitar duplicar "en media tensión" si ya viene en solicitud)
    ubic = (meta or {}).get("ubicacion") or ""
    solicitud = (meta or {}).get("solicitud") or "Factibilidad y Punto de Diseño"
    if ubic:
        sol_l = solicitud.lower()
        if "media tensión" in sol_l or "media tension" in sol_l:
            mid = "evaluación de %s, predio ubicado en %s." % (solicitud, ubic)
        else:
            mid = "evaluación de %s en media tensión, predio ubicado en %s." % (solicitud, ubic)
        for p in root.iter("{%s}p" % W_NS):
            txt = _para_text(p)
            if "Unidad de Proyectos" in txt and "predio ubicado" in txt.lower():
                new_txt = "La Unidad de Proyectos y Obras distribución solicita %s" % mid
                if _set_para_text(p, new_txt):
                    done.append({"old": "antecedentes ubic…", "new": new_txt[:80], "count": 1})
                break
    return done


def _xml_escape_text(value):
    """Escape minimo para contenido de <w:t>."""
    s = "" if value is None else str(value)
    return (
        s.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def _iter_wt_texts(root):
    """Lista de textos de cada <w:t> en orden de documento."""
    out = []
    for t in root.iter("{%s}t" % W_NS):
        out.append(t.text if t.text is not None else "")
    return out


def _patch_document_xml_wt(xml_bytes, new_texts):
    """
    Reescribe solo el contenido de <w:t>…</w:t> en orden, sin reserializar el XML.
    ElementTree.tostring/write corrompe namespaces de Word; este parche los preserva.
    """
    pattern = re.compile(br"(<w:t(?:\s[^>]*)?>)(.*?)(</w:t>)", re.DOTALL)
    matches = list(pattern.finditer(xml_bytes))
    if len(matches) != len(new_texts):
        raise ValueError(
            "Conteo <w:t> distinto tras edicion: xml=%d tree=%d"
            % (len(matches), len(new_texts))
        )
    parts = []
    last = 0
    changed = 0
    for i, m in enumerate(matches):
        parts.append(xml_bytes[last:m.start()])
        old_inner = m.group(2)
        new_inner = _xml_escape_text(new_texts[i]).encode("utf-8")
        if old_inner != new_inner:
            changed += 1
        parts.append(m.group(1) + new_inner + m.group(3))
        last = m.end()
    parts.append(xml_bytes[last:])
    return b"".join(parts), changed


def fill_word(docx_path, scenarios, meta, images_dir=None, excel_tables=None):
    tmp = docx_path + ".__fill_tmp"
    if os.path.isdir(tmp):
        shutil.rmtree(tmp)
    mkdir(tmp)
    with zipfile.ZipFile(docx_path, "r") as zin:
        zin.extractall(tmp)

    word_dir = os.path.join(tmp, "word")
    # document.xml + headers (REFERENCIA EMAPICA) — nunca reserializar con ET.write
    xml_targets = ["document.xml"]
    if os.path.isdir(word_dir):
        for name in sorted(os.listdir(word_dir)):
            low = name.lower()
            if low.startswith("header") and low.endswith(".xml"):
                xml_targets.append(name)

    header_done = []
    tables_report = {"n_tables": 0, "filled": [], "cells": []}
    wt_changed_total = 0

    for xml_name in xml_targets:
        xml_path = os.path.join(word_dir, xml_name)
        if not os.path.isfile(xml_path):
            continue
        with open(xml_path, "rb") as f:
            original_xml = f.read()
        tree = ET.parse(xml_path)
        root = tree.getroot()
        done = _fill_word_headers(root, meta, scenarios)
        if xml_name == "document.xml":
            tables_report = _fill_word_tables(root, scenarios, meta, excel_tables=excel_tables)
            header_done = done
        elif done:
            header_done.extend([dict(x, where=xml_name) for x in done])
        new_texts = _iter_wt_texts(root)
        patched_xml, wt_changed = _patch_document_xml_wt(original_xml, new_texts)
        with open(xml_path, "wb") as f:
            f.write(patched_xml)
        wt_changed_total += wt_changed

    images_replaced = []
    if images_dir and os.path.isdir(images_dir):
        for src_name, dst_rel in IMAGE_MAP.items():
            src = os.path.join(images_dir, src_name)
            dst = os.path.join(tmp, dst_rel.replace("/", os.sep))
            if os.path.isfile(src) and os.path.isfile(dst):
                shutil.copy2(src, dst)
                images_replaced.append({"src": src, "dst": dst_rel})

    out_tmp = docx_path + ".__new.docx"
    if os.path.isfile(out_tmp):
        os.remove(out_tmp)
    with zipfile.ZipFile(out_tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for folder, _, files in os.walk(tmp):
            for name in files:
                full = os.path.join(folder, name)
                rel = os.path.relpath(full, tmp).replace("\\", "/")
                zout.write(full, rel)
    shutil.move(out_tmp, docx_path)
    shutil.rmtree(tmp, ignore_errors=True)
    return {
        "replacements": header_done,
        "tables_filled": tables_report.get("filled") or [],
        "tables_cells": tables_report.get("cells") or [],
        "tables_source": tables_report.get("source"),
        "n_tables": tables_report.get("n_tables"),
        "images_replaced": images_replaced,
        "wt_patched": wt_changed_total,
        "xml_targets": xml_targets,
    }


def delivery_status(settings=None):
    """Estado de gates de entrega sin rellenar Word/Excel ni regenerar charts."""
    s = settings or load_settings()
    scenarios = _load_scenarios(s)
    meta = _meta_cliente(s)
    ocr = load_informe_meta(s) or {}
    img_dir = _images_dir(s)
    mkdir(img_dir)
    present = [f for f in REQUIRED_LF_IMAGES if os.path.isfile(os.path.join(img_dir, f))]
    missing_imgs = [f for f in REQUIRED_LF_IMAGES if f not in present]
    ready, missing = _check_delivery_gates(scenarios, meta, img_dir, charts_res=None, settings=s)
    return {
        "ok": True,
        "delivery_ready": ready,
        "missing": missing,
        "feeder_id": s.get("feeder_id"),
        "checks": {
            "loadflow_situacional": bool(scenarios.get("situacional")) and _lf_status_ok(scenarios.get("raw_situacional")),
            "loadflow_proyectado": bool(scenarios.get("proyectado")) and _lf_status_ok(scenarios.get("raw_proyectado")),
            "informe_meta_ocr": meta_is_complete(ocr),
            "lf_images": len(missing_imgs) == 0,
        },
        "meta": {
            "cliente": meta.get("cliente") or "",
            "potencia_kw": meta.get("potencia_kw"),
            "potencia_txt": meta.get("potencia_txt") or "",
            "alimentador": meta.get("alimentador") or "",
            "meta_source": meta.get("meta_source"),
            "complete_ocr": meta_is_complete(ocr),
        },
        "images_dir": img_dir,
        "images_present": present,
        "images_missing": missing_imgs,
        "required_lf_images": list(REQUIRED_LF_IMAGES),
        "scenarios": {
            "situacional": bool(scenarios.get("situacional")),
            "proyectado": bool(scenarios.get("proyectado")),
        },
    }


def _metric_snapshot(scenarios):
    metric_keys = ("kw", "kvar", "kva", "fp_pct", "kw_loss", "vpu", "v_pct_a", "v_pct_b", "v_pct_c", "i_a")
    return {
        "situacional": bool(scenarios.get("situacional")),
        "proyectado": bool(scenarios.get("proyectado")),
        "situacional_metrics": (
            {k: scenarios["situacional"].get(k) for k in metric_keys}
            if scenarios.get("situacional") else None
        ),
        "proyectado_metrics": (
            {k: scenarios["proyectado"].get(k) for k in metric_keys}
            if scenarios.get("proyectado") else None
        ),
    }


def _confirm_path(paths):
    return os.path.join(paths["doc_dir"], "entrega_confirmada.json")


def load_entrega_confirmada(settings=None):
    """Lee confirmación de entrega (cierre tras vista preliminar)."""
    s = settings or load_settings()
    paths = informe_paths(s)
    path = _confirm_path(paths)
    if not os.path.isfile(path):
        return None
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return None


def build_informe_preview(settings=None):
    """
    Vista preliminar del informe ya creado (sin regenerar).
    Usar antes de cerrar/confirmar entrega.
    """
    s = settings or load_settings()
    scenarios = _load_scenarios(s)
    meta = _meta_cliente(s)
    paths = scenarios["paths"]
    img_dir = _images_dir(s)
    mkdir(img_dir)

    images = []
    for name in list(REQUIRED_LF_IMAGES) + ["topologia.png", "trafo_cargabilidad.png"]:
        full = os.path.join(img_dir, name)
        if os.path.isfile(full):
            images.append({
                "name": name,
                "url": "/api/informe/imagen/%s" % name,
                "bytes": os.path.getsize(full),
                "mtime": datetime.fromtimestamp(os.path.getmtime(full)).isoformat(timespec="seconds"),
                "required": name in REQUIRED_LF_IMAGES,
            })

    doc_inf = paths.get("informe_doc") or ""
    doc_jus = paths.get("justificacion_doc") or ""
    doc_pdf = os.path.join(paths["doc_dir"], "informe.pdf") if paths.get("doc_dir") else ""
    docs = {
        "informe": {
            "exists": bool(doc_inf and os.path.isfile(doc_inf)),
            "path": doc_inf,
            "url": "/api/informe/archivo/informe",
            "bytes": os.path.getsize(doc_inf) if doc_inf and os.path.isfile(doc_inf) else 0,
        },
        "justificacion": {
            "exists": bool(doc_jus and os.path.isfile(doc_jus)),
            "path": doc_jus,
            "url": "/api/informe/archivo/justificacion",
            "bytes": os.path.getsize(doc_jus) if doc_jus and os.path.isfile(doc_jus) else 0,
        },
        "pdf": {
            "exists": bool(doc_pdf and os.path.isfile(doc_pdf)),
            "path": doc_pdf,
            "url": "/api/informe/archivo/pdf",
            "bytes": os.path.getsize(doc_pdf) if doc_pdf and os.path.isfile(doc_pdf) else 0,
        },
    }

    # Paginas rasterizadas del render Word (si existen)
    page_previews = []
    prev_dir = os.path.join(paths["doc_dir"], "informe_preview") if paths.get("doc_dir") else ""
    if prev_dir and os.path.isdir(prev_dir):
        for name in sorted(os.listdir(prev_dir)):
            if not name.lower().endswith((".jpg", ".jpeg", ".png")):
                continue
            full = os.path.join(prev_dir, name)
            page_previews.append({
                "name": name,
                "url": "/api/informe/pagina/%s" % name,
                "bytes": os.path.getsize(full),
            })

    filled = docs["informe"]["exists"] and docs["justificacion"]["exists"]
    required_imgs = [i for i in images if i.get("required")]
    preview_ok = filled and len(required_imgs) >= len(REQUIRED_LF_IMAGES)
    confirm = load_entrega_confirmada(s)

    word_replacements = []
    man_path = os.path.join(paths["doc_dir"], "fill_manifest.json")
    if os.path.isfile(man_path):
        try:
            with open(man_path, "r", encoding="utf-8") as f:
                man = json.load(f)
            word_replacements = ((man.get("word") or {}).get("replacements") or [])[:40]
            if not scenarios.get("situacional") and man.get("scenarios_used"):
                # fallback metrics from last fill
                pass
        except Exception:
            man = {}
    else:
        man = {}

    scenarios_used = man.get("scenarios_used") or _metric_snapshot(scenarios)
    render_info = man.get("render") or {}

    return {
        "ok": True,
        "preview_ok": preview_ok,
        "can_close": preview_ok and not bool((confirm or {}).get("confirmed")),
        "closed": bool((confirm or {}).get("confirmed")),
        "confirm": confirm,
        "feeder_id": s.get("feeder_id"),
        "filled_at": man.get("filled_at"),
        "meta": {
            "cliente": meta.get("cliente") or "",
            "ubicacion": meta.get("ubicacion") or "",
            "solicitud": meta.get("solicitud") or "",
            "potencia_kw": meta.get("potencia_kw"),
            "potencia_txt": meta.get("potencia_txt") or "",
            "alimentador": meta.get("alimentador") or s.get("feeder_id") or "",
            "set": meta.get("set") or "",
            "tension_kv": meta.get("tension_kv"),
            "transformador": meta.get("transformador") or "",
            "expediente": meta.get("expediente") or "",
            "proyecto": meta.get("proyecto") or "",
            "meta_source": meta.get("meta_source"),
        },
        "scenarios_used": scenarios_used,
        "images": images,
        "images_dir": img_dir,
        "docs": docs,
        "page_previews": page_previews,
        "render": {
            "ok": bool(render_info.get("ok")),
            "pages": render_info.get("pages"),
            "pdf": render_info.get("pdf") or (doc_pdf if docs["pdf"]["exists"] else None),
            "rendered_at": render_info.get("rendered_at"),
        },
        "word_replacements": word_replacements,
        "msg": (
            "Vista preliminar lista — valide datos y gráficas antes de cerrar."
            if preview_ok else
            "Informe incompleto: rellene §6.2 (Word/Excel + 4 gráficas LF)."
        ),
    }


def confirm_informe_entrega(settings=None, note="", force=False):
    """
    Cierra la entrega tras validar la vista preliminar.
    Exige docs + 4 PNG; escribe doc/entrega_confirmada.json.
    """
    s = settings or load_settings()
    preview = build_informe_preview(s)
    if not preview.get("preview_ok") and not force:
        return {
            "ok": False,
            "error": "No se puede cerrar: falta vista preliminar completa (docs o gráficas).",
            "preview": preview,
        }
    paths = informe_paths(s)
    mkdir(paths["doc_dir"])
    payload = {
        "ok": True,
        "confirmed": True,
        "confirmed_at": datetime.now().isoformat(timespec="seconds"),
        "feeder_id": s.get("feeder_id"),
        "note": (note or "").strip(),
        "preview_snapshot": {
            "meta": preview.get("meta"),
            "scenarios_used": preview.get("scenarios_used"),
            "images": [i.get("name") for i in (preview.get("images") or [])],
            "docs": {
                "informe": (preview.get("docs") or {}).get("informe", {}).get("path"),
                "justificacion": (preview.get("docs") or {}).get("justificacion", {}).get("path"),
            },
            "filled_at": preview.get("filled_at"),
        },
    }
    path = _confirm_path(paths)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, ensure_ascii=False, default=str)
    payload["path"] = path
    payload["msg"] = "Entrega confirmada · informe validado en vista preliminar"
    return payload


def fill_informe(settings=None, overwrite_copy=True, require_delivery=True):
    """
    Copia plantilla → doc/ y rellena Excel + Word con LoadFlow.

    require_delivery=True (default): exige ambos LF, meta OCR minima y 4 PNG LF.
    Si faltan, no toca doc/ y retorna ok=False + missing[].
    """
    s = settings or load_settings()
    scenarios = _load_scenarios(s)
    meta = _meta_cliente(s)
    # Persistir meta completa (SET/trafo/etc.) para UI §6.1
    try:
        from pipeline.extract_informe_meta_pdf import save_informe_meta
        save_informe_meta(s, meta)
    except Exception as ex:
        notes = ["aviso save meta: %s" % ex]
    else:
        notes = []
    paths = scenarios["paths"]
    img_dir = _images_dir(s)
    mkdir(img_dir)
    charts_res = None

    # Preferir capturas CYMDIST vivas (API: coloreo + ExportActiveView / GUI).
    # Orden: 1) captura API situacional+proyectado  2) mapa inventario (fallback)
    #        3) graficas matplotlib solo si aun falta slot
    capture_res = None
    prefer_cymdist = bool(s.get("informe_prefer_cymdist_captures", True))
    # Por defecto SI: integrar capturas estado actual / con proyecto via CYMDIST API
    auto_capture = s.get("informe_auto_cymdist_capture")
    if auto_capture is None:
        auto_capture = True
    else:
        auto_capture = bool(auto_capture)
    force_cap = bool(s.get("force_cymdist_captures", False))
    force_charts = bool(s.get("force_informe_charts", False))
    if prefer_cymdist:
        try:
            from pipeline.capture_informe_color_views import (
                is_cymdist_capture,
                is_live_cymdist_view,
                ensure_standard_legends,
                capture_informe_color_views,
            )
            from pipeline.render_informe_color_maps import generate_informe_color_maps
            ensure_standard_legends(force=False)

            missing_live = [
                f for f in REQUIRED_LF_IMAGES
                if not is_live_cymdist_view(os.path.join(img_dir, f))
            ]
            if auto_capture and (missing_live or force_cap):
                notes.append(
                    "captura CYMDIST API: activando coloreo VoltageLevel/LoadingLevel "
                    "en situacional + proyectado…"
                )
                capture_res = capture_informe_color_views(
                    settings=s,
                    open_gui=bool(s.get("informe_capture_open_gui", True)),
                    force=force_cap or bool(missing_live),
                )
                # Persist log for UI diagnostics
                try:
                    with open(os.path.join(img_dir, "capture_log.txt"), "a", encoding="utf-8") as lf:
                        lf.write("\n--- fill_informe capture %s ---\n" % datetime.now().isoformat(timespec="seconds"))
                        lf.write(json.dumps(capture_res, indent=2, ensure_ascii=False, default=str))
                        lf.write("\n")
                except Exception:
                    pass
                notes.append(
                    "cymdist API captures: %d ok, err: %d"
                    % (
                        len(capture_res.get("generated") or []),
                        len(capture_res.get("errors") or []),
                    )
                )
                for err in (capture_res.get("errors") or [])[:4]:
                    notes.append("capture: %s" % err)
            else:
                notes.append("cymdist live views ya presentes (4 PNG)")

            # Fallback inventario solo para slots que NO tienen captura viva
            still_missing = [
                f for f in REQUIRED_LF_IMAGES
                if not is_live_cymdist_view(os.path.join(img_dir, f))
                and not is_cymdist_capture(os.path.join(img_dir, f))
            ]
            if still_missing or (
                force_cap and not (capture_res and capture_res.get("generated"))
            ):
                rend = generate_informe_color_maps(
                    settings=s, force=False  # nunca pisar live views
                )
                notes.append(
                    "color maps fallback: %d gen, %d skip, err: %d"
                    % (
                        len(rend.get("generated") or []),
                        len(rend.get("skipped") or []),
                        len(rend.get("errors") or []),
                    )
                )
                for err in (rend.get("errors") or [])[:3]:
                    notes.append("color map: %s" % err)
            else:
                notes.append("sin fallback topology_render (capturas API OK)")
        except Exception as ex:
            notes.append("cymdist color omitido: %s" % ex)

    if scenarios.get("situacional") or scenarios.get("proyectado"):
        try:
            charts_res = generate_informe_charts(
                img_dir,
                {
                    "situacional": scenarios.get("situacional"),
                    "proyectado": scenarios.get("proyectado"),
                },
                paths=paths,
                force=force_charts,
            )
            notes.append(
                "charts generated: %d skipped: %d"
                % (len(charts_res.get("generated") or []), len(charts_res.get("skipped") or []))
            )
        except Exception as ex:
            charts_res = {
                "ok": False,
                "errors": [str(ex)],
                "generated": [],
                "missing": list(REQUIRED_LF_IMAGES),
            }
            notes.append("charts error: %s" % ex)

    # Mapa satelite §2.1: ubicacion de la carga nueva (nodo X/Y → WGS84)
    map_res = None
    try:
        from pipeline.generate_location_map import generate_location_map
        map_res = generate_location_map(s, out_dir=img_dir, force=True)
        if map_res.get("ok"):
            notes.append("topologia.png mapa carga nueva OK")
        else:
            notes.append("topologia.png omitido: %s" % (map_res.get("error") or "?"))
    except Exception as ex:
        map_res = {"ok": False, "error": str(ex)}
        notes.append("topologia.png error: %s" % ex)

    delivery_ready, missing = _check_delivery_gates(
        scenarios, meta, img_dir, charts_res, settings=s
    )

    scenarios_used = _metric_snapshot(scenarios)

    if require_delivery and not delivery_ready:
        err = (
            "Informe incompleto para entrega. Falta: %s. "
            "Complete: PDF OCR en §5 (cliente+potencia), "
            "Flujo situacional, Flujo proyectado; las graficas LF se generan solas."
            % (", ".join(missing) if missing else "requisitos")
        )
        manifest = {
            "ok": False,
            "delivery_ready": False,
            "error": err,
            "missing": missing,
            "filled_at": datetime.now().isoformat(timespec="seconds"),
            "feeder_id": s.get("feeder_id"),
            "meta": meta,
            "meta_source": meta.get("meta_source"),
            "paths": paths,
            "images_dir": img_dir,
            "image_slots": IMAGE_MAP,
            "required_lf_images": list(REQUIRED_LF_IMAGES),
            "charts_generated": (charts_res or {}).get("generated") or [],
            "charts": charts_res,
            "cymdist_captures": capture_res,
            "location_map": map_res,
            "scenarios_used": scenarios_used,
            "notes": notes,
            "aviso_imagenes": (
                "Graficas LF se generan desde loadflow_*.json. "
                "topologia.png = mapa satelite carga nueva. Override en %s."
                % img_dir
            ),
        }
        man_path = os.path.join(paths["doc_dir"], "fill_manifest.json")
        mkdir(paths["doc_dir"])
        with open(man_path, "w", encoding="utf-8") as f:
            json.dump(manifest, f, indent=2, ensure_ascii=False, default=str)
        manifest["fill_manifest"] = man_path
        print("[fill] INCOMPLETO:", err)
        return manifest

    # Solo si entrega lista: copiar plantilla y rellenar
    base = assemble_informe(s, overwrite=overwrite_copy)
    if not base.get("ok"):
        return base

    # Rellenar de nuevo invalida cierre previo (hay que revalidar preliminar).
    try:
        conf = _confirm_path(paths)
        if os.path.isfile(conf):
            os.remove(conf)
            notes.append("entrega_confirmada.json invalidada (relleno nuevo)")
    except Exception as ex:
        notes.append("aviso invalidar confirm: %s" % ex)

    xnotes, excel_tables = fill_excel(paths["justificacion_doc"], scenarios, meta)
    notes.extend(xnotes)
    wres = fill_word(
        paths["informe_doc"], scenarios, meta,
        images_dir=img_dir, excel_tables=excel_tables,
    )
    notes.append("word replacements: %d" % len(wres.get("replacements") or []))
    _tbl = ", ".join(wres.get("tables_filled") or []) or "(none)"
    notes.append("word tables: %s" % _tbl)
    notes.append("word tables source: %s" % (wres.get("tables_source") or excel_tables and "excel" or "metrics"))
    notes.append("images replaced: %d" % len(wres.get("images_replaced") or []))

    replaced_names = [
        os.path.basename(x.get("src") or "") for x in (wres.get("images_replaced") or [])
    ]
    lf_replaced = [n for n in REQUIRED_LF_IMAGES if n in replaced_names]

    # Render final Word: campos, paginado y PDF de entrega
    render_res = None
    try:
        from pipeline.render_informe import render_informe as _render_informe
        render_res = _render_informe(s, export_pdf=True, preview_pages=True)
        if render_res.get("ok"):
            notes.append(
                "render Word ok · paginas=%s · pdf=%s"
                % (render_res.get("pages"), "si" if render_res.get("pdf") else "no")
            )
            for n in (render_res.get("notes") or [])[:8]:
                notes.append("render: %s" % n)
        else:
            notes.append("aviso render Word: %s" % (render_res.get("error") or "fallo"))
    except Exception as ex:
        render_res = {"ok": False, "error": str(ex)}
        notes.append("aviso render Word: %s" % ex)

    manifest = {
        "ok": True,
        "delivery_ready": True,
        "missing": [],
        "filled_at": datetime.now().isoformat(timespec="seconds"),
        "feeder_id": s.get("feeder_id"),
        "meta": meta,
        "meta_source": meta.get("meta_source"),
        "paths": paths,
        "images_dir": img_dir,
        "image_slots": IMAGE_MAP,
        "required_lf_images": list(REQUIRED_LF_IMAGES),
        "charts_generated": (charts_res or {}).get("generated") or [],
        "charts": charts_res,
        "cymdist_captures": capture_res,
        "location_map": map_res,
        "scenarios_used": scenarios_used,
        "excel_notes": xnotes,
        "excel_tables": {
            "mapping": (excel_tables or {}).get("mapping"),
            "situacional": (excel_tables or {}).get("situacional"),
            "proyectado": (excel_tables or {}).get("proyectado"),
        },
        "word": wres,
        "render": render_res,
        "lf_images_replaced": lf_replaced,
        "notes": notes,
        "preview_url": "/api/informe/preview",
        "needs_preview_confirm": True,
        "aviso_imagenes": (
            "Graficas LF desde JSON en %s. "
            "topologia.png = mapa satelite de la carga nueva (nodo de conexion). "
            "Valide la vista preliminar (§6.3) antes de cerrar la entrega."
            % img_dir
        ),
    }
    man_path = os.path.join(paths["doc_dir"], "fill_manifest.json")
    with open(man_path, "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2, ensure_ascii=False, default=str)
    manifest["fill_manifest"] = man_path
    print("[fill] Excel+Word actualizados en", paths["doc_dir"])
    return manifest


def main():
    m = fill_informe()
    print(json.dumps(m, indent=2, ensure_ascii=False, default=str))
    if not m.get("ok"):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
