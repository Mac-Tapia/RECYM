# -*- coding: utf-8 -*-
"""§3.3b · Excel de verificación de distribución de carga."""
import os
import sys
import tempfile
import unittest

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(ROOT, "src"))

from pipeline import distribution_report as dr


def _load(load_id, kw, kvar, kwh, kva=250.0, locked=False, connected=True):
    return {"LoadID": load_id, "SectionID": "", "kW": kw, "kvar": kvar, "kWh": kwh,
            "kVA_instalado": kva, "locked": locked, "connected": connected}


class TestDistributionReport(unittest.TestCase):
    def setUp(self):
        self.loads = [
            _load("DEV_2010_1_SE100", 100.0, 32.87, 10000.0),
            _load("DEV_2010_2_SE200", 300.0, 98.61, 30000.0, kva=400.0),
            _load("DEV_2010_3_SE300", 500.0, 164.34, 1.0, kva=630.0, locked=True),
            _load("DEV_2010_4_SE400", 0.0, 0.0, 0.0),
        ]
        self.clientes = [{"Cliente": "Exalmar", "LoadID_CYMDIST": "DEV_2010_3_SE300",
                          "Pot": 500.0, "Activo": True}]

    def _by_sed(self, rows):
        return {r["SED"]: r for r in rows}

    def test_columns_and_links(self):
        rows, summary = dr.build_distribution_report(
            self.loads, self.clientes, cabecera={"P_kW": 920.0, "Q_kvar": 300.0})
        self.assertEqual(list(rows[0].keys()), dr.COLUMNS)
        by = self._by_sed(rows)
        self.assertEqual(by["SE300"]["accion"], dr.ACCION_CI)
        self.assertIn("cliente: Exalmar", by["SE300"]["observaciones"])
        self.assertEqual(by["SE300"]["Kw"], 500.0)
        self.assertEqual(by["SE100"]["accion"], dr.ACCION_DISTRIBUIDO)
        # 10 000 de 40 000 kWh → 25 % de los 400 kW distribuidos
        self.assertAlmostEqual(by["SE100"]["Kw"], 100.0, places=3)
        self.assertTrue(by["SE100"]["observaciones"].startswith("OK"))
        self.assertEqual(by["SE400"]["accion"], dr.ACCION_SIN_CONSUMO)
        # Toda SED debe tener carga: sin consumo en la BD se marca para revisar.
        self.assertTrue(by["SE400"]["observaciones"].startswith("REVISAR"))
        self.assertEqual(summary["sed_sin_consumo"], ["SE400"])
        self.assertEqual(summary["diferencia_cabecera_kW"], 20.0)
        self.assertEqual(summary["n_revisar"], 1)

    def test_flags_unlocked_client_wrong_share_and_overload(self):
        loads = [dict(l) for l in self.loads]
        loads[2]["locked"] = False
        loads[2]["kW"] = 450.0
        loads[0]["kW"] = 160.0          # debería ser ~25 % del distribuido
        loads[0]["kVA_instalado"] = 150.0
        rows, summary = dr.build_distribution_report(loads, self.clientes)
        by = self._by_sed(rows)
        self.assertTrue(by["SE300"]["observaciones"].startswith("REVISAR"))
        self.assertIn("sin Locked", by["SE300"]["observaciones"])
        self.assertIn("Δkw", by["SE100"]["observaciones"])
        self.assertIn("sobrecarga", by["SE100"]["observaciones"])
        self.assertGreaterEqual(summary["n_revisar"], 2)

    def test_excluded_client_expected_zero(self):
        clientes = [dict(self.clientes[0], Activo=False)]
        loads = [dict(l) for l in self.loads]
        loads[2]["kW"] = loads[2]["kvar"] = 0.0
        loads[2]["connected"] = False
        rows, _ = dr.build_distribution_report(loads, clientes)
        row = self._by_sed(rows)["SE300"]
        self.assertEqual(row["accion"], dr.ACCION_EXCLUIDO)
        self.assertEqual(row["Kw"], 0.0)
        self.assertTrue(row["observaciones"].startswith("OK"))

    def test_sed_and_numbers(self):
        self.assertEqual(dr.sed_from_load_id("DEV_2010_329293_SE20131"), "SE20131")
        self.assertEqual(dr.sed_from_load_id("DEV_2010_1180302_M-21251"), "M-21251")
        self.assertEqual(dr.sed_from_load_id("DEV_2010_330974_M20063"), "M20063")
        self.assertEqual(dr.num("248,80301"), 248.80301)
        self.assertEqual(dr.num("1.234,5"), 1234.5)

    def test_xlsx_has_exact_header_and_summary(self):
        from openpyxl import load_workbook

        rows, summary = dr.build_distribution_report(self.loads, self.clientes, {"P_kW": 920.0})
        path = os.path.join(tempfile.mkdtemp(prefix="recym_dist_"), "d.xlsx")
        dr.write_distribution_xlsx(path, rows, summary, feeder_id="TM105")
        wb = load_workbook(path)
        header = [c.value for c in wb["Distribucion"][1]]
        self.assertEqual(header, dr.COLUMNS)
        self.assertEqual(wb["Distribucion"].max_row, len(rows) + 1)
        self.assertEqual(wb["Resumen"]["B1"].value, "TM105")



class TestAllocationKeepsFixedClients(unittest.TestCase):
    def test_locked_clients_are_not_released_by_default(self):
        from core.cymdist_com import allocation_unlock_loads_flag

        self.assertEqual(allocation_unlock_loads_flag({}), 0)
        self.assertEqual(allocation_unlock_loads_flag(None), 0)
        self.assertEqual(allocation_unlock_loads_flag({"loadallocation_unlock_loads": True}), 1)



class TestMeterTwins(unittest.TestCase):
    def test_meter_with_loaded_sed_is_ok_and_empty_sed_is_flagged(self):
        loads = [
            _load("DEV_2010_1_SE20646", 32.85, 8.0, 7295.0),
            _load("DEV_2010_2_M20646", 0.0, 0.0, 0.0, kva=15.0),
            _load("DEV_2010_3_SE20923", 0.0, 0.0, 0.0),
            _load("DEV_2010_4_M20923", 0.0, 0.0, 0.0, kva=15.0),
        ]
        rows, summary = dr.build_distribution_report(loads, [])
        by = {r["SED"]: r for r in rows}
        self.assertEqual(by["M20646"]["accion"], dr.ACCION_MEDIDOR)
        self.assertTrue(by["M20646"]["observaciones"].startswith("OK"))
        self.assertIn("carga en SE20646", by["M20646"]["observaciones"])
        self.assertTrue(by["SE20923"]["observaciones"].startswith("REVISAR"))
        self.assertIn("KWHUsage", by["SE20923"]["observaciones"])
        self.assertIn("tampoco tiene consumo", by["M20923"]["observaciones"])
        self.assertEqual(summary["sed_sin_consumo"], ["SE20923"])
        self.assertEqual(dr.meter_twin_sed("M-21251"), "SE21251")


class TestDigsilentExport(unittest.TestCase):
    def test_balanced_rows_by_sed_with_units(self):
        loads = [
            _load("DEV_2010_1_SE100", 100.0, 30.0, 10.0),
            _load("DEV_2010_9_SE100", 50.0, 10.0, 5.0),
            _load("DEV_2010_3_SE300", 500.0, 164.34, 1.0, connected=False),
        ]
        clientes = [{"Cliente": "Exalmar", "LoadID_CYMDIST": "DEV_2010_1_SE100", "Pot": 100.0}]
        rows = {r["SED"]: r for r in dr.digsilent_rows(loads, clientes, feeder_id="TM105")}
        self.assertEqual(list(rows["SE100"].keys()), dr.DIGSILENT_COLUMNS)
        self.assertEqual(rows["SE100"]["P_kW"], 150.0)
        self.assertEqual(rows["SE100"]["Q_kvar"], 40.0)
        self.assertEqual(rows["SE100"]["P_MW"], 0.15)
        self.assertEqual(rows["SE100"]["Cliente_importante"], "Exalmar")
        self.assertEqual(rows["SE300"]["P_kW"], 0.0)
        self.assertEqual(rows["SE300"]["Conectada"], "No")

    def test_workbook_includes_digsilent_sheet(self):
        from openpyxl import load_workbook

        loads = [_load("DEV_2010_1_SE100", 100.0, 30.0, 10.0)]
        rows, summary = dr.build_distribution_report(loads, [])
        path = os.path.join(tempfile.mkdtemp(prefix="recym_dg_"), "d.xlsx")
        dr.write_distribution_xlsx(path, rows, summary, "TM105", digsilent=dr.digsilent_rows(loads, []))
        sheet = load_workbook(path)["DIgSILENT"]
        self.assertEqual([c.value for c in sheet[1]], dr.DIGSILENT_COLUMNS)
        self.assertEqual(sheet["A2"].value, "SE100")

    def test_inventory_row_is_three_phase_and_keeps_kvar(self):
        from pipeline.inventory_loads import inventory_row

        load = dict(_load("DEV_2010_7_SE700", 248.803, 65.801, 68479.4, kva=160.0),
                    Tipo="SpotLoad", n_values=3)
        row = inventory_row(load, "NET_TM105")
        self.assertEqual(row["SED"], "SE700")
        self.assertEqual(row["kW"], "248.803")
        self.assertEqual(row["kvar"], "65.801")
        self.assertAlmostEqual(row["FP"], 0.9668, places=4)
        self.assertEqual(row["n_values"], 3)


if __name__ == "__main__":
    unittest.main()
